"""Helpers and the 9 October 2026 templates. Every element is a native PowerPoint object.
Sizes are real points on an 11.25 in wide slide (1080 px = 810 pt), as in the spec. Positions are px."""
import math, os, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_PATTERN
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "carousels"))
from photos import crop, ROOT   # noqa: E402

PX = 9525
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
DEEP, GREEN, CREAM, GLOW, LIGHT, BROWN, INK, GOLD = "1E3F1D", "31662F", "F3F1EA", "B1BCB1", "B1BCB1", "4C3634", "333130", "D39C48"
PHBG, STRIPE, WHITE, FAINT = "C9C4B8", "E8E2D4", "FFFFFF", "6A665C"
HEAD, BODY, TAG = "Montserrat", "Source Sans 3", "[COPY: Allison]"
REGISTRY = []   # every PHOTO placeholder: (file, slide number, label)

def rgb(h): return RGBColor.from_string(h)

# ---------------------------------------------------------------- text estimation
def est_lines(text, pt, width, bold=False):
    cw = pt * 1.3333 * (0.64 if bold else 0.50)
    n = 0
    for para in text.split("\n"):
        line = 0; n += 1
        for word in para.split(" "):
            w = len(word) * cw
            if line and line + cw + w > width: n += 1; line = w
            else: line += (cw if line else 0) + w
    return n

def fit(text, width, sizes, maxlines, bold=False):
    for s in sizes:
        if est_lines(text, s, width, bold) <= maxlines: return s
    return sizes[-1]

def th(text, pt, width, bold=False, spacing=1.2):
    return est_lines(text, pt, width, bold) * pt * 1.3333 * spacing

# ---------------------------------------------------------------- deck
class Deck:
    def __init__(self, w=1080, h=1350, name=""):
        self.w, self.h, self.name = w, h, name
        self.prs = Presentation(); self.prs.slide_width = Emu(w*PX); self.prs.slide_height = Emu(h*PX)
        self.meta = []

    def slide(self, bg=CREAM, note="", counter=True, tid=""):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        if bg:
            s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(bg)
        s.notes_slide.notes_text_frame.text = (f"{tid}. " if tid else "") + note
        self.meta.append(dict(dark=bg in (DEEP, GREEN, BROWN) if bg else False, counter=counter, bg=bg))
        s._deck = self
        return s

    def finish(self, path, counters=True):
        n = len(self.prs.slides)
        if counters and n > 1:
            for i, (s, m) in enumerate(zip(self.prs.slides, self.meta), 1):
                if m["counter"]:
                    counter(s, i, n, m["dark"], self.w)
        if os.path.dirname(path): os.makedirs(os.path.dirname(path), exist_ok=True)
        self.prs.save(path)

def slide_no(s):
    return [sl.slide_id for sl in s._deck.prs.slides].index(s.slide_id) + 1

# ---------------------------------------------------------------- placement (with optional tilt about a centre)
class Tilt:
    def __init__(self, cx, cy, deg): self.cx, self.cy, self.deg = cx, cy, deg

def _place(sh, x, y, w, h, tl=None):
    if tl is None:
        sh.left, sh.top, sh.width, sh.height = Emu(int(x*PX)), Emu(int(y*PX)), Emu(int(w*PX)), Emu(int(h*PX)); return
    a = math.radians(tl.deg); dx, dy = x + w/2 - tl.cx, y + h/2 - tl.cy
    nx = tl.cx + dx*math.cos(a) - dy*math.sin(a); ny = tl.cy + dx*math.sin(a) + dy*math.cos(a)
    sh.left, sh.top, sh.width, sh.height = Emu(int((nx-w/2)*PX)), Emu(int((ny-h/2)*PX)), Emu(int(w*PX)), Emu(int(h*PX))
    sh.rotation = tl.deg

def _alpha(fill_parent, alpha):
    clr = fill_parent.find(".//{%s}srgbClr" % A)
    for old in clr.findall("{%s}alpha" % A): clr.remove(old)
    e = etree.SubElement(clr, "{%s}alpha" % A); e.set("val", str(int(alpha*1000)))

# ---------------------------------------------------------------- shapes
def rect(s, x, y, w, h, fill=None, alpha=None, line=None, lw=1, kind=MSO_SHAPE.RECTANGLE, radius=None, tl=None, shadow=False, pattern=None):
    r = s.shapes.add_shape(kind, 0, 0, Emu(1), Emu(1))
    if pattern:
        r.fill.patterned(); r.fill.pattern = MSO_PATTERN.LIGHT_UPWARD_DIAGONAL
        r.fill.fore_color.rgb = rgb(pattern); r.fill.back_color.rgb = rgb(fill)
    elif fill:
        r.fill.solid(); r.fill.fore_color.rgb = rgb(fill)
        if alpha is not None: _alpha(r._element.spPr.find("{%s}solidFill" % A), alpha)
    else: r.fill.background()
    if line: r.line.color.rgb = rgb(line); r.line.width = Emu(int(lw*PX))
    else: r.line.fill.background()
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        r.adjustments[0] = min(0.5, radius / min(w, h))
    r.shadow.inherit = False
    if shadow:
        sp = r._element.spPr
        for e in sp.findall("{%s}effectLst" % A): sp.remove(e)
        eff = etree.SubElement(sp, "{%s}effectLst" % A)
        sh = etree.SubElement(eff, "{%s}outerShdw" % A, blurRad="228600", dist="76200", dir="5400000", algn="t", rotWithShape="0")
        c = etree.SubElement(sh, "{%s}srgbClr" % A, val="000000"); etree.SubElement(c, "{%s}alpha" % A, val="30000")
    _place(r, x, y, w, h, tl)
    return r

def oval(s, x, y, w, h, fill=None, **kw): return rect(s, x, y, w, h, fill, kind=MSO_SHAPE.OVAL, **kw)
def rrect(s, x, y, w, h, fill=None, radius=24, **kw): return rect(s, x, y, w, h, fill, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius, **kw)

def poly(s, pts, fill, tl=None, line=None, lw=1):
    xs, ys = [p[0] for p in pts], [p[1] for p in pts]
    fb = s.shapes.build_freeform(int(pts[0][0]*PX), int(pts[0][1]*PX), scale=1.0)
    fb.add_line_segments([(int(px_*PX), int(py_*PX)) for px_, py_ in pts[1:]], close=True)
    sh = fb.convert_to_shape()
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    if line: sh.line.color.rgb = rgb(line); sh.line.width = Emu(int(lw*PX))
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    _place(sh, min(xs), min(ys), max(xs)-min(xs), max(ys)-min(ys), tl)
    return sh

def gradient(s, x, y, w, h, color, a0, a1):
    r = rect(s, x, y, w, h, color)
    sp = r._element.spPr
    for tag in ("solidFill", "noFill", "gradFill"):
        for e in sp.findall("{%s}%s" % (A, tag)): sp.remove(e)
    g = etree.Element("{%s}gradFill" % A, rotWithShape="1"); lst = etree.SubElement(g, "{%s}gsLst" % A)
    for pos, al in ((0, a0), (100000, a1)):
        gs = etree.SubElement(lst, "{%s}gs" % A, pos=str(pos)); c = etree.SubElement(gs, "{%s}srgbClr" % A, val=color)
        etree.SubElement(c, "{%s}alpha" % A, val=str(int(al*1000)))
    etree.SubElement(g, "{%s}lin" % A, ang="5400000", scaled="0")
    sp.insert(list(sp).index(sp.find("{%s}prstGeom" % A)) + 1, g)
    return r

# ---------------------------------------------------------------- text
def text(s, x, y, w, h, t, pt, color=INK, font=BODY, bold=False, italic=False, align="l", anchor="t",
         spacing=1.2, alpha=None, track=None, outline=None, tl=None):
    tb = s.shapes.add_textbox(0, 0, Emu(1), Emu(1)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = dict(t=MSO_ANCHOR.TOP, m=MSO_ANCHOR.MIDDLE, b=MSO_ANCHOR.BOTTOM)[anchor]
    for i, line in enumerate(t.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = dict(l=PP_ALIGN.LEFT, c=PP_ALIGN.CENTER, r=PP_ALIGN.RIGHT)[align]; p.line_spacing = spacing
        r = p.add_run(); r.text = line; f = r.font
        f.size = Pt(pt); f.name = font; f.bold = bold; f.italic = italic; f.color.rgb = rgb(color)
        rPr = r._r.get_or_add_rPr()
        if alpha is not None: _alpha(rPr.find("{%s}solidFill" % A), alpha)
        if track: rPr.set("spc", str(int(track*100)))
        if outline:
            ln = etree.Element("{%s}ln" % A, w=str(int(outline[1]*PX))); sf = etree.SubElement(ln, "{%s}solidFill" % A)
            etree.SubElement(sf, "{%s}srgbClr" % A, val=outline[0]); rPr.insert(0, ln)
    _place(tb, x, y, w, h, tl)
    return tb

def label(s, x, y, w, t, pt=22, color=GREEN, **kw):
    return text(s, x, y, w, pt*1.33*1.3, t.upper(), pt, color, HEAD, True, track=2, **kw)

def pill(s, x, y, t, pt, fg, bg, alpha=None, w=None, h=None, pad=16, font=HEAD, track=None, tl=None):
    tw = w or (len(t) * pt * 1.3333 * (0.74 if track else 0.68) + 2*pad)
    hh = h or (pt * 1.3333 * 1.25 + 2*pad)
    rrect(s, x, y, tw, hh, bg, alpha=alpha, radius=min(40, hh/2), tl=tl)
    text(s, x, y, tw, hh, t, pt, fg, font, True, align="c", anchor="m", track=track, tl=tl)
    return tw, hh

def counter(s, i, n, dark, W=1080):
    text(s, W - 80 - 140, 80, 140, 30, f"{i}/{n}", 20,
         CREAM if dark else INK, BODY, align="r", alpha=60)

# ---------------------------------------------------------------- photos and marks
MENTOR_DIR = os.path.join(ROOT, "assets", "mentors-teachers-day")
MENTOR_KEYS = [("carla", "carla-portugal"), ("nick padwick", "nick-padwick"), ("wes sander", "wes-sander"),
               ("gerald", "gerald-ramirez"), ("caterina", "caterina-capri")]

def mentor_photo(label):
    """A portrait label such as 'Dr. Carla Portugal, square headshot' -> file in assets/mentors-teachers-day/, if uploaded."""
    l = label.lower()
    if "portrait" not in l and "headshot" not in l: return None
    for key, slug in MENTOR_KEYS:
        if key in l:
            for ext in ("jpg", "jpeg", "png"):
                f = os.path.join(MENTOR_DIR, f"{slug}.{ext}")
                if os.path.exists(f): return os.path.relpath(f, ROOT)
    return None

def photo(s, x, y, w, h, src, note="", radius=0, kind=None, fx=0.5, fy=0.5, deckname=""):
    src = mentor_photo(src) or src
    """src: repo path (str starting 'assets/') -> cropped picture; otherwise a PHOTO placeholder label."""
    if src and src.startswith("assets/"):
        p = crop(src, int(w), int(h), fx, fy)
        pic = s.shapes.add_picture(p, Emu(int(x*PX)), Emu(int(y*PX)), Emu(int(w*PX)), Emu(int(h*PX)))
        if kind == "oval": pic.auto_shape_type = MSO_SHAPE.OVAL
        elif radius:
            pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
            av = pic._element.spPr.find("{%s}prstGeom" % A).find("{%s}avLst" % A)
            if av is None: av = etree.SubElement(pic._element.spPr.find("{%s}prstGeom" % A), "{%s}avLst" % A)
            etree.SubElement(av, "{%s}gd" % A, name="adj", fmla="val %d" % int(min(50000, radius/min(w, h)*100000)))
        if w * h > 0.6 * s._deck.w * s._deck.h: s._deck.meta[slide_no(s) - 1]["dark"] = True
        return pic
    k = MSO_SHAPE.OVAL if kind == "oval" else (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    r = rect(s, x, y, w, h, PHBG, kind=k, radius=radius or None)
    small = min(w, h) < 300
    text(s, x + 16, y, w - 32, h, f"PHOTO: {src}", 20 if not small else 16, "4A463F", BODY, True, align="c", anchor="m")
    REGISTRY.append((s._deck.name, slide_no(s), src))
    return r

LOGO_W = os.path.join(ROOT, "assets", "logo", "foundation-logo-white.png")
LOGO_C = os.path.join(ROOT, "assets", "logo", "foundation-logo-color.png")
LOGO_AR = 743 / 668

def logo(s, x=None, y=None, size=90, outline=CREAM, tl=None, cx=None):
    """Soil Food Web Foundation logo: white version on dark backgrounds (outline==CREAM), colour version on light ones.
    size = logo height in px. cx centres it on that x."""
    w = s._deck.w; h = s._deck.h; pw = size * LOGO_AR
    if cx is not None: x = cx - pw / 2
    x = w - 80 - pw if x is None else x; y = h - 80 - size if y is None else y
    pic = s.shapes.add_picture(LOGO_W if outline == CREAM else LOGO_C, 0, 0, Emu(int(pw * PX)), Emu(int(size * PX)))
    _place(pic, x, y, pw, size, tl)
    return pic

def logo_for(s, dark=True, **kw): logo(s, outline=CREAM if dark else DEEP, **kw)

def series_pill(s, y=80):
    pill(s, 80, y, "SOIL REGENERATORS IN THE WILD", 22, GLOW, DEEP, alpha=80, w=760, h=70, track=2)

def quote_mark(s, x, y, color, alpha, pt=260):
    text(s, x, y, 300, pt*1.333, "“", pt, color, HEAD, True, alpha=alpha, spacing=0.9)

def check(s, cx, cy, d, color=GREEN):
    oval(s, cx-d/2, cy-d/2, d, d, color)
    k = d/140
    poly(s, [(cx-38*k, cy+2*k), (cx-24*k, cy-12*k), (cx-8*k, cy+4*k), (cx+30*k, cy-34*k), (cx+44*k, cy-20*k), (cx-8*k, cy+32*k)], CREAM)

def warn(s, cx, cy, d, color=GOLD, fg=DEEP):
    oval(s, cx-d/2, cy-d/2, d, d, color)
    text(s, cx-d/2, cy-d/2, d, d, "!", 60 * d/140, fg, HEAD, True, align="c", anchor="m", spacing=1.0)

def follow_pill(s, cx, y, t, pt=28, maxw=880):
    lines = est_lines(t, pt, maxw - 64, True)
    w = min(maxw, len(t) * pt * 1.3333 * 0.68 + 64) if lines == 1 else maxw
    h = lines * pt * 1.3333 * 1.25 + 40
    rrect(s, cx - w/2, y, w, h, GREEN, radius=min(40, h/2))
    text(s, cx - w/2 + 32, y, w - 64, h, t, pt, CREAM, HEAD, True, align="c", anchor="m")
    return h

# ================================================================ TEMPLATE 1: Soil Regenerators in the wild
def t1a(d, photo_src, name, place, note="", tid="1A"):
    s = d.slide(DEEP, note or "Series cover. Same cover on every Soil Regenerators post.", tid=tid)
    photo(s, 0, 0, 1080, 1350, photo_src)
    gradient(s, 0, 743, 1080, 607, DEEP, 0, 85)
    series_pill(s)
    one = est_lines(name, 64, 800, True) == 1
    ny = 1090 if one else 990
    rect(s, 80, ny - 34, 120, 4, LIGHT)
    nh = th(name, 64, 800, True, 1.1)
    text(s, 80, ny, 800, nh + 6, name, 64, CREAM, HEAD, True, spacing=1.1, anchor="t")
    text(s, 80, ny + nh + 18, 780, 44, place, 30, GLOW, BODY, spacing=1.1)
    logo(s)
    d.meta[-1]["dark"] = True
    return s

def t1b(d, quote, who, note="", tid="1B"):
    s = d.slide(CREAM, note, tid=tid)
    quote_mark(s, 80, 120, GREEN, 30)
    sz = fit(quote, 880, [48, 44, 40], 7, True)
    h = th(quote, sz, 880, True, 1.2)
    top = max(330, 640 - h/2)
    text(s, 100, top, 880, h + 10, quote, sz, DEEP, HEAD, True, align="c", spacing=1.2)
    text(s, 100, top + h + 40, 880, 40, who, 26, GREEN, BODY, align="c")
    return s

def t1c(d, photo_src, lab, line, note="", tid="1C", **pk):
    s = d.slide(CREAM, note, tid=tid)
    photo(s, 0, 0, 1080, 840, photo_src, **pk)
    label(s, 80, 890, 920, lab, 20, GREEN)
    sz = fit(line, 920, [38, 36, 34], 3, False)
    text(s, 80, 945, 920, 300, line, sz, DEEP, BODY, spacing=1.3)
    return s

def t1d(d, lab, line, note="", tid="1D", big=None, sub=None):
    s = d.slide(DEEP, note, tid=tid)
    if big:
        sz = fit(big, 880, [96, 88, 80], 3, True); h = th(big, sz, 880, True, 1.1)
        sub_h = th(sub, 34, 880, False, 1.3) if sub else 0
        top = 675 - (h + sub_h + 40) / 2
        label(s, 80, top - 60, 920, lab, 20, GLOW)
        text(s, 80, top, 880, h + 10, big, sz, CREAM, HEAD, True, spacing=1.1)
        if sub: text(s, 80, top + h + 40, 880, sub_h + 10, sub, 34, GLOW, BODY, spacing=1.3)
        return s
    sz = fit(line, 880, [52, 48, 44], 7, True); h = th(line, sz, 880, True, 1.15)
    top = 675 - h / 2
    label(s, 80, top - 70, 920, lab, 20, GLOW)
    text(s, 80, top, 880, h + 10, line, sz, CREAM, HEAD, True, spacing=1.15)
    return s

def t1e(d, quote, who, follow, lead=None, note="", tid="1E"):
    s = d.slide(CREAM, note, tid=tid)
    if quote:
        quote_mark(s, 80, 120, GREEN, 30)
        sz = fit(quote, 880, [48, 44, 40], 6, True); h = th(quote, sz, 880, True, 1.2)
        top = max(330, 560 - h/2)
        text(s, 100, top, 880, h + 10, quote, sz, DEEP, HEAD, True, align="c", spacing=1.2)
        y = top + h + 40
        if who: text(s, 100, y, 880, 40, who, 26, GREEN, BODY, align="c"); y += 90
        follow_pill(s, 540, y, follow)
    else:
        if lead:
            text(s, 100, 480, 880, 200, lead, 48, DEEP, HEAD, True, align="c", spacing=1.2)
        follow_pill(s, 540, 640 if lead else 560, follow, 30 if lead else 40)
    logo(s, outline=DEEP)
    return s

# ================================================================ TEMPLATE 2: Did you know / myth-buster
def t2a(d, photo_src, lab, headline, small=None, strike=False, gold_rule=False, note="", tid="2A", logo_on=True):
    s = d.slide(DEEP, note, tid=tid)
    photo(s, 0, 0, 1080, 640, photo_src)
    rect(s, 0, 640, 1080, 8, GLOW)
    label(s, 80, 700, 800, lab, 22, GLOW)
    sz = fit(headline, 800, [64, 60, 56], 4, True); h = th(headline, sz, 800, True, 1.1)
    if gold_rule: rect(s, 80, 745, 120, 5, GOLD)
    text(s, 80, 770, 800, h + 10, headline, sz, CREAM, HEAD, True, spacing=1.1)
    if strike: rect(s, 60, 770 + h/2 - 6, 840, 12, GLOW)
    if small: text(s, 80, 770 + h + 25, 800, 45, small, 30, GLOW, BODY, spacing=1.2)
    if logo_on: logo(s)
    return s

def t2b(d, num, line, img=None, extra=None, note="", tid="2B"):
    s = d.slide(CREAM, note, tid=tid)
    if num is not None: text(s, 80, 70, 400, 190, str(num), 140, GREEN, HEAD, True, alpha=25, spacing=1.0)
    sz = fit(line, 880, [46, 42, 40], 6, True); h = th(line, sz, 880, True, 1.1)
    text(s, 100, 330 if num is None else 420, 880, h + 10, line, sz, DEEP, HEAD, True, spacing=1.1)
    if extra:
        y = (330 if num is None else 420) + h + 40
        for e in extra: text(s, 100, y, 880, 60, e, 34, INK, BODY, spacing=1.3); y += 70
    if img is not None:
        photo(s, 600, 870, 400, 400, img, radius=24, fx=0.42)
    return s

def t2c(d, q, note="", tid="2C"):
    s = d.slide(CREAM, note, tid=tid)
    text(s, 520, 300, 520, 500, "?", 300, LIGHT, HEAD, True, align="c", alpha=20, spacing=1.0)
    sz = fit(q, 640, [50, 46, 42], 8, True); h = th(q, sz, 640, True, 1.2)
    text(s, 100, 675 - h/2, 640, h + 10, q, sz, DEEP, HEAD, True, spacing=1.2)
    return s

def icon_save(s, cx, y):
    poly(s, [(cx-30, y), (cx+30, y), (cx+30, y+80), (cx, y+58), (cx-30, y+80)], CREAM)
    poly(s, [(cx-22, y+8), (cx+22, y+8), (cx+22, y+64), (cx, y+48), (cx-22, y+64)], GREEN)

def icon_send(s, cx, y):
    poly(s, [(cx-44, y+36), (cx+44, y), (cx+10, y+80), (cx-4, y+46)], CREAM)
    poly(s, [(cx-34, y+36), (cx+30, y+10), (cx-4, y+40)], GREEN)

def t2d(d, line, extra=None, note="", tid="2D"):
    s = d.slide(GREEN, note, tid=tid)
    sz = fit(line, 880, [54, 50, 46], 6, True); h = th(line, sz, 880, True, 1.2) * 1.15
    top = 480 - h/2
    text(s, 100, top, 880, h + 10, line, sz, CREAM, HEAD, True, align="c", spacing=1.2)
    y = top + h + 30
    if extra: text(s, 100, y, 880, 50, extra, 34, GLOW, BODY, align="c"); y += 80
    y = max(y + 40, 900)
    icon_save(s, 380, y); text(s, 280, y + 100, 200, 40, "Save this", 22, CREAM, HEAD, True, align="c")
    icon_send(s, 700, y); text(s, 600, y + 100, 200, 40, "Send this", 22, CREAM, HEAD, True, align="c")
    logo(s)
    return s

# ================================================================ TEMPLATE 3: Checklist / yes-no
def t3a(d, headline, photos, tag="SAVE THIS", note="", tid="3A"):
    s = d.slide(CREAM, note, tid=tid)
    pill(s, 80, 80, tag, 20, CREAM, GREEN, track=2, h=60)
    sz = fit(headline, 920, [64, 60, 56], 5, True)
    text(s, 80, 220, 920, th(headline, sz, 920, True, 1.1) + 10, headline, sz, DEEP, HEAD, True, spacing=1.1)
    if len(photos) == 1:
        photo(s, 80, 850, 920, 420, photos[0], radius=24)
    else:
        for i, p in enumerate(photos):
            photo(s, 80 + i * (220 + 13.33), 1050, 220, 220, p, radius=24)
    return s

def t3b(d, item, expl, verdicts=("check",), img=None, note="", tid="3B"):
    s = d.slide(CREAM, note, tid=tid)
    photo(s, 160, 110, 760, 760, img, radius=32)
    for i, v in enumerate(verdicts):
        cx = 920 - 20 - 70 - i * 160; cy = 110 + 20 + 70 - 0
        (check(s, cx, cy, 140) if v == "check" else warn(s, cx, cy, 140, GOLD if v == "warn" else BROWN, DEEP if v == "warn" else CREAM))
    isz = fit(item, 920, [56, 50, 46], 2, True); ih = th(item, isz, 920, True, 1.1)
    text(s, 80, 895, 920, ih + 6, item, isz, DEEP, HEAD, True, spacing=1.1)
    sz = fit(expl, 920, [34, 32, 30], 4, False)
    text(s, 80, 895 + ih + 22, 920, 1270 - (895 + ih + 22), expl, sz, INK, BODY, spacing=1.3)
    return s

def t3c(d, head, items, note="", tid="3C"):
    s = d.slide(CREAM, note, tid=tid)
    label(s, 80, 180, 900, head, 22, GREEN)
    for i, it in enumerate(items):
        y = 300 + i * 200
        oval(s, 80, y, 80, 80, GREEN)
        text(s, 80, y, 80, 80, str(i + 1), 34, CREAM, HEAD, True, align="c", anchor="m", spacing=1.0)
        text(s, 200, y - 10, 780, 110, it, 38, DEEP, BODY, spacing=1.3, anchor="m")
    return s

# ================================================================ TEMPLATE 4: Mentor
def scallop(s, cx, cy, R, tl):
    for k in range(16):
        a = 2 * math.pi * k / 16; px_, py_ = cx + math.cos(a) * (R - 12), cy + math.sin(a) * (R - 12)
        oval(s, px_ - 24, py_ - 24, 48, 48, GREEN, tl=tl)
    for k in range(16):
        a = 2 * math.pi * k / 16; px_, py_ = cx + math.cos(a) * (R - 12), cy + math.sin(a) * (R - 12)
        oval(s, px_ - 21, py_ - 21, 42, 42, CREAM, tl=tl)
    oval(s, cx - R + 14, cy - R + 14, 2 * R - 28, 2 * R - 28, CREAM, tl=tl)

def ribbon(s, x, y, w, h, fill, tl, notch=30):
    return poly(s, [(x, y), (x + w, y), (x + w - notch, y + h/2), (x + w, y + h), (x, y + h), (x + notch, y + h/2)], fill, tl=tl)

def t4a(d, photo_src, name, role, note="", tid="4A"):
    s = d.slide(DEEP, note, counter=True, tid=tid)
    tl = Tilt(540, 675, 3)
    cx0, cy0 = 100, 85
    rrect(s, cx0, cy0, 880, 1180, CREAM, radius=28, tl=tl, shadow=True, pattern=STRIPE)
    rect(s, cx0 + 60 - 10, cy0 + 60 - 10, 400, 400, CREAM, line=DEEP, lw=2, tl=tl)
    photo_tilt(s, cx0 + 60, cy0 + 60, 380, photo_src, tl)
    scallop(s, cx0 + 880 - 60 - 90, cy0 + 60 + 90, 90, tl)
    bx_, by_ = cx0 + 880 - 60 - 90, cy0 + 60 + 90
    pic = s.shapes.add_picture(LOGO_C, 0, 0, Emu(int(110 * PX)), Emu(int(110 / LOGO_AR * PX))); _place(pic, bx_ - 55, by_ - 55 / LOGO_AR, 110, 110 / LOGO_AR, tl)
    rows = [("BASED IN", ""), ("TEACHING SINCE", ""), ("ASK ME ABOUT", ""), ("FAVORITE ORGANISM", "")]
    for i, (lab, _) in enumerate(rows):
        y = cy0 + 60 + 400 + 20 + i * 100
        text(s, cx0 + 60, y, 760, 30, lab, 20, GREEN, HEAD, True, track=2, tl=tl)
        text(s, cx0 + 60, y + 30, 760, 44, TAG, 30, INK, BODY, tl=tl)
    ribbon(s, cx0 + 60, 1000, 760, 110, GREEN, tl)
    nsz = fit(name.upper(), 620, [50, 46, 42, 38], 1, True)
    text(s, cx0 + 60 + 40, 1000, 680, 110, name.upper(), nsz, CREAM, HEAD, True, align="c", anchor="m", tl=tl)
    ribbon(s, cx0 + 60 + 60, 1000 + 70 + 30, 640, 80, DEEP, tl)
    text(s, cx0 + 60 + 60 + 40, 1100, 560, 80, role, 28, GLOW, BODY, align="c", anchor="m", tl=tl)
    rect(s, cx0, 85 + 1180 - 44, 880, 44, BROWN, tl=tl)
    text(s, cx0, 85 + 1180 - 44, 880, 44, "soilfoodweb.com", 18, CREAM, BODY, align="c", anchor="m", tl=tl)
    return s

def photo_tilt(s, x, y, size, src, tl):
    src = mentor_photo(src) or src
    if src and src.startswith("assets/"):
        p = crop(src, size, size)
        pic = s.shapes.add_picture(p, 0, 0, Emu(size*PX), Emu(size*PX)); _place(pic, x, y, size, size, tl); return
    r = rect(s, x, y, size, size, PHBG, tl=tl)
    text(s, x + 16, y, size - 32, size, f"PHOTO: {src}", 20, "4A463F", BODY, True, align="c", anchor="m", tl=tl)
    REGISTRY.append((s._deck.name, slide_no(s), src))

def t4b(d, headline, note="", tid="4B"):
    s = d.slide(DEEP, note, tid=tid)
    sz = fit(headline, 880, [64, 60, 56], 4, True)
    text(s, 80, 250, 880, th(headline, sz, 880, True, 1.1) + 10, headline, sz, CREAM, HEAD, True, spacing=1.1)
    for dx, deg in ((-140, -6), (0, 0), (140, 6)):
        cx, cy = 440 + dx, 1000
        rrect(s, cx - 260, cy - 350, 520, 700, CREAM, radius=28, tl=Tilt(cx, cy, deg), shadow=True, pattern=STRIPE)
    return s

def t4c(d, photo_src, name, role, lab="MEET OUR MENTORS", quote=None, duo=None, note="", tid="4C", role2=None, name2=None):
    s = d.slide(DEEP, note, counter=True, tid=tid)
    if duo:
        photo(s, 0, 0, 537, 1350, duo[0]); photo(s, 543, 0, 537, 1350, duo[1])
        rect(s, 537, 0, 6, 1350, CREAM)
    else:
        photo(s, 0, 0, 1080, 1350, photo_src)
    pill(s, 80, 80, lab, 22, GLOW, DEEP, track=2, h=70, alpha=90)
    if quote:
        gradient(s, 0, 640, 1080, 350, DEEP, 0, 70)
        quote_mark(s, 80, 660, GLOW, 100, 120)
        text(s, 80, 780, 920, 200, quote, 40, CREAM, HEAD, True, spacing=1.2, anchor="b")
    rect(s, 0, 990, 1080, 360, DEEP, alpha=90)
    if name2:
        for x0, n, r_ in ((80, name, role), (580, name2, role2)):
            sz = fit(n, 420, [44, 40, 36], 2, True)
            text(s, x0, 1040, 420, 150, n, sz, CREAM, HEAD, True, spacing=1.1)
            text(s, x0, 1190, 420, 70, r_, 28, GLOW, BODY, spacing=1.2)
        return s
    else:
        sz = fit(name, 800, [58, 50, 44, 40, 36], 3, True)
        text(s, 80, 1030, 800, th(name, sz, 800, True, 1.1) + 10, name, sz, CREAM, HEAD, True, spacing=1.1)
        if role: text(s, 80, 1030 + th(name, sz, 800, True, 1.1) + 20, 800, 50, role, 30, GLOW, BODY, spacing=1.2)
    logo(s)
    return s

# ================================================================ TEMPLATE 5: Big number
def t5(d, num, units, context, note="", tid="5", bars=None, src=True):
    s = d.slide(CREAM, note, tid=tid)
    if bars:
        base = 1090
        for i, (n, lab, hh, col) in enumerate(bars):
            x = 180 + i * 460
            rect(s, x, base - hh, 260, hh, col)
            text(s, x - 40, base - hh - 118, 340, 100, n, 72, col, HEAD, True, align="c", anchor="b", spacing=1.0)
            text(s, x - 40, base + 24, 340, 90, lab, 30, INK, BODY, align="c", spacing=1.2)
        rect(s, 120, base, 840, 3, DEEP)
        if context: text(s, 100, 130, 880, 150, context, 32, INK, BODY, align="c", spacing=1.3)
        if src: text(s, 80, 1230, 920, 40, "Source: " + TAG, 24, FAINT, BODY)
        return s
    if isinstance(num, tuple): text(s, 100, 330, 880, 40, num[1], 24, GREEN, HEAD, True, align="c")
    n = num[0] if isinstance(num, tuple) else num
    text(s, 40, 400, 1000, 280, n, 200, GREEN, HEAD, True, align="c", anchor="m", spacing=1.0)
    text(s, 100, 700, 880, 160, units, 56, DEEP, HEAD, True, align="c", spacing=1.1)
    text(s, 140, 700 + th(units, 56, 880, True, 1.1) + 30, 800, 120, context, 32, INK, BODY, align="c", spacing=1.3)
    if src: text(s, 80, 1230, 920, 40, "Source: " + TAG, 24, FAINT, BODY)
    return s

# ================================================================ TEMPLATE 6: Reel cover and caption cards (1080 x 1920)
def t6a(d, photo_src, title, series=False, split=None, circle=False, note="", tid="6A"):
    if circle:
        s = d.slide(DEEP, note, counter=False, tid=tid)
        photo(s, 140, 300, 800, 800, photo_src, kind="oval")
        sz = fit(title, 880, [72, 66, 60], 3, True); h = th(title, sz, 880, True, 1.1)
        text(s, 100, 1180, 880, h + 10, title, sz, CREAM, HEAD, True, align="c", spacing=1.1)
        return s
    s = d.slide(DEEP, note, counter=False, tid=tid)
    if split:
        photo(s, 0, 0, 540, 1920, split[0]); photo(s, 540, 0, 540, 1920, split[1])
        pill(s, 100, 400, "FIELD", 26, GLOW, DEEP, track=2, alpha=90); pill(s, 640, 400, "COMPOST", 26, GLOW, DEEP, track=2, alpha=90)
    else:
        photo(s, 0, 0, 1080, 1920, photo_src)
    if series: series_pill(s, 300)
    sz = fit(title, 800, [84, 76, 70], 3, True); h = th(title, sz, 800, True, 1.1)
    top = 900 - h/2
    rrect(s, 80, top - 32, 920, h + 64, DEEP, alpha=85, radius=20)
    text(s, 120, top, 840, h + 4, title, sz, CREAM, HEAD, True, align="c", spacing=1.1)
    return s

def t6b(d, line, note="", tid="6B"):
    s = d.slide(None, "Transparent card: remove the page background in Canva (or export as PNG with transparent background). " + note, counter=False, tid=tid)
    sz = fit(line, 880, [64, 58, 52], 3, True); h = th(line, sz, 880, True, 1.1)
    text(s, 100, 1150 - h/2, 880, h + 10, line, sz, CREAM, HEAD, True, align="c", spacing=1.1, outline=(DEEP, 6))
    return s

# ================================================================ TEMPLATE 7: Story
def t7(d, photo_src, headline, line, note="", tid="7"):
    s = d.slide(DEEP, note, counter=False, tid=tid)
    photo(s, 0, 0, 1080, 1920, photo_src)
    rect(s, 0, 0, 1080, 1920, DEEP, alpha=50)
    sz = fit(headline, 880, [72, 66, 60], 5, True); h = th(headline, sz, 880, True, 1.1)
    text(s, 100, 420, 880, h + 10, headline, sz, CREAM, HEAD, True, align="c", spacing=1.1)
    text(s, 100, 420 + h + 90, 880, 150, line, 44, GLOW, BODY, align="c", spacing=1.3)
    logo(s, cx=540, y=1500)
    return s

# ================================================================ TEMPLATE 8: Donate
def t8_donate(d, line, note="", tid="8"):
    s = d.slide(DEEP, note, tid=tid)
    sz = fit(line, 880, [56, 52, 48], 5, True); h = th(line, sz, 880, True, 1.2)
    text(s, 100, 520 - h/2, 880, h + 10, line, sz, CREAM, HEAD, True, align="c", spacing=1.2)
    rrect(s, 280, 800, 520, 110, GOLD, radius=40)
    text(s, 280, 800, 520, 110, "Donate: link in bio", 30, DEEP, HEAD, True, align="c", anchor="m")
    logo(s)
    return s

# ================================================================ TEMPLATE 9: LinkedIn
def t9(d, photo_src, note="", tid="9"):
    s = d.slide(DEEP, note, counter=False, tid=tid)
    photo(s, 0, 0, d.w, d.h, photo_src)
    logo(s, x=d.w - 50 - 70 * LOGO_AR, y=d.h - 50 - 70, size=70)
    return s

def t9_wes(d, left, right, note="", tid="9"):
    s = d.slide(DEEP, note, counter=False, tid=tid)
    photo(s, 0, 0, 540, 1080, left); photo(s, 540, 0, 540, 1080, right)
    pill(s, 35, 800, "Conventional field", 26, CREAM, DEEP, alpha=90, w=470, h=110); pill(s, 575, 800, "Biologically active compost", 26, CREAM, DEEP, alpha=90, w=470, h=110)
    logo(s, x=1080 - 50 - 70 * LOGO_AR, y=1080 - 50 - 70, size=70)
    return s
