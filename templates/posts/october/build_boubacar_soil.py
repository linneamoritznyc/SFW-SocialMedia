"""Boubacar seamless carousel, soil cross-section design (no colour fades).
Cream paper above, one continuous brand-brown soil layer below with groundcover and roots crossing every seam.
Same fixed grid as the hyphae version: photos in the same area on every slide, text in the same block in the soil.
Brand colours and fonts only.

python3 tools/make_soil_panorama.py assets/collage/boubacar-soil 11
python3 build_boubacar_soil.py  ->  01-10-2026-thu-ig-soil-regenerators-boubacar-soil.pptx
"""
import os
from pptx.util import Emu
from lib import Deck, rect, text, crop, ROOT, PX
from build_oct_01_10 import logo, save
from build_boubacar import B, SOURCES
from build_boubacar_panorama import tw, nlines, lh, HEADF, BODYF, VOICE, PHOTO_Y, PHOTO_H, MAIN_X, MAIN_W, COL_X, COL_W

W, H = 1080, 1350
PANO = "assets/collage/boubacar-soil/slide-{:02d}.jpg"
CREAM, BROWN, GREEN, TAN, SAGE = "F3F1EA", "4C3634", "31662F", "C09D7F", "B1BCB1"
TEXT_X, TEXT_W, TEXT_Y, TEXT_H = 80, 920, 900, 330     # text block in the soil, identical on every slide


def slide(d, i, note):
    s = d.slide(None, note, counter=False)
    s.shapes.add_picture(os.path.join(ROOT, PANO.format(i)), 0, 0, Emu(W * PX), Emu(H * PX))
    return s


def pic(s, rel, x, y, w, h, fx=0.5, fy=0.5):
    rect(s, x - 8, y - 8, w + 16, h + 16, "FFFFFF", shadow=True)
    s.shapes.add_picture(crop(rel, int(w), int(h), fx, fy), Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))


def main(s, rel, fx=0.5, fy=0.5): pic(s, rel, MAIN_X, PHOTO_Y, MAIN_W, PHOTO_H - 30, fx, fy)
def col1(s, rel, fx=0.5, fy=0.5): pic(s, rel, COL_X, PHOTO_Y, COL_W, PHOTO_H - 30, fx, fy)


def col2(s, a, b, fa=(0.5, 0.5), fb=(0.5, 0.5)):
    h = (PHOTO_H - 30 - 36) / 2
    pic(s, a, COL_X, PHOTO_Y, COL_W, h, *fa); pic(s, b, COL_X, PHOTO_Y + h + 36, COL_W, h, *fb)


def words(s, body, head=None, voice=False):
    """Text set straight onto the soil: cream heading, sage body, cream italic for his words."""
    bfont = VOICE if voice else BODYF
    for hpt, bpt in ((52, 54), (50, 50), (48, 46), (46, 42), (44, 40), (42, 38), (40, 36), (38, 34), (36, 32), (34, 30)):
        hh = nlines(head, hpt, TEXT_W * 0.96, HEADF) * lh(hpt, 1.05) + 16 if head else 0
        bh = nlines(body, bpt, TEXT_W * 0.96, bfont) * lh(bpt, 1.15)
        if hh + bh <= TEXT_H: break
    y = TEXT_Y + (TEXT_H - hh - bh) / 2
    if head:
        text(s, TEXT_X, y, TEXT_W, hh, head, hpt, CREAM, HEADF, True, spacing=1.05); y += hh
    text(s, TEXT_X, y, TEXT_W, bh + 10, body, bpt, CREAM if voice else SAGE, bfont, False, italic=voice, spacing=1.15)


def quote_mark(s):
    text(s, COL_X + 10, PHOTO_Y - 30, COL_W, 300, "“", 260, GREEN, VOICE, False, spacing=0.8)


def footer(s, t="soilfoodweb.com"):
    text(s, 0, H - 92, W, 50, t, 24, CREAM, HEADF, True, align="c", anchor="m")


def build():
    d = Deck(name="01-10-2026-thu-ig-soil-regenerators-boubacar-soil")
    N = (" Soil cross-section seamless carousel: background assets/collage/boubacar-soil/ (tools/make_soil_panorama.py, drawn by code, "
         "brand colours). His quotes from his email to Allison, 30 Sep 2026.")

    s = slide(d, 1, "Cover." + N + " " + SOURCES + " Graduation line from Linnea: confirm course and date [VERIFY].")
    main(s, B + "boubacar-portrait-bananas-closer.jpg", 0.5, 0.35)
    logo(s, COL_X + 20, PHOTO_Y + 10, 240, white=False)
    text(s, COL_X, PHOTO_Y + 300, COL_W, 300, "Happy\nInternational\nCoffee Day", 28, GREEN, HEADF, True, spacing=1.15)
    words(s, "Guinea-Conakry\nGraduated from the Soil Food Web School in July 2026", "Boubacar Tidiane Diallo")

    s = slide(d, 2, "The land. NEW COPY from the caption draft." + N + " Photos: farm-overview-with-tanks.jpeg, young-tree-with-pineapples.jpg, man-by-water-tanks.jpeg.")
    main(s, B + "farm-overview-with-tanks.jpeg")
    col2(s, B + "young-tree-with-pineapples.jpg", B + "man-by-water-tanks.jpeg", (0.5, 0.55), (0.35, 0.5))
    words(s, "The soil at Gnaly Coffee & AgroÉcole Bio, in the Fouta Djallon highlands of Guinea-Conakry, had been degraded "
             "and poorly managed for years.", "When he started, the land was worn out.")
    footer(s)

    s = slide(d, 3, "The question." + N + " Photo: vegetable-beds-by-building.jpeg.")
    main(s, B + "vegetable-beds-by-building.jpeg", 0.6, 0.25)
    quote_mark(s)
    words(s, "“Today, I ask a different question: What does the soil food web need to thrive?”", "So he changed the question:", True)
    footer(s)

    s = slide(d, 4, "Feeding the soil." + N + " Photos: boubacar-compost-pile.jpeg; still-biochar-kiln-smoke-73s.png (making-biochar-in-barrel-kiln.mp4 at 1:13).")
    main(s, B + "boubacar-compost-pile.jpeg", 0.5, 0.45)
    col1(s, B + "still-biochar-kiln-smoke-73s.png", 0.5, 0.4)
    words(s, "living groundcover, mulch from his farm and the forest, and compost with biochar activated with fish hydrolysate "
             "and compost tea.", "How he feeds his soil:")
    footer(s)

    s = slide(d, 5, "In practice." + N + " Photos: carrying-mulch.jpg, still-biochar-charcoal-16s.png, barrel-brew-under-shelter.jpeg.")
    main(s, B + "carrying-mulch.jpg")
    col2(s, B + "still-biochar-charcoal-16s.png", B + "barrel-brew-under-shelter.jpeg", (0.5, 0.5), (0.5, 0.35))
    words(s, "“I apply compost regularly around young trees and whenever I add new mulch.”", voice=True)
    footer(s)

    s = slide(d, 6, "What grows." + N + " Photos: planting-seedling-in-agroforest.jpeg, banana-bunch.jpg.")
    main(s, B + "planting-seedling-in-agroforest.jpeg", 0.45, 0.45)
    col1(s, B + "banana-bunch.jpg", 0.5, 0.4)
    words(s, "coffee, cacao, bananas, citrus, and pineapple.", "What he grows together:")
    footer(s)

    s = slide(d, 7, "His compost." + N + " Photos: buckets-of-tubers.jpg, young-tree-yellow-new-leaves.jpg, man-with-harvested-tubers.jpeg [VERIFY who].")
    main(s, B + "buckets-of-tubers.jpg")
    col2(s, B + "young-tree-yellow-new-leaves.jpg", B + "man-with-harvested-tubers.jpeg", (0.5, 0.45), (0.5, 0.5))
    words(s, "“My compost includes plant residues, animal manure when available, biochar activated with fish hydrolysate and "
             "compost tea, and local organic materials.”", voice=True)
    footer(s)

    s = slide(d, 8, "What's changing." + N + " Photos: millipede-leaf-litter.jpg, boubacar-hand-compost-worm.jpg.")
    main(s, B + "millipede-leaf-litter.jpg", 0.5, 0.35)
    col1(s, B + "boubacar-hand-compost-worm.jpg", 0.4, 0.5)
    words(s, "more earthworms, more mushrooms, active decomposition, and better soil structure.", "What he's seeing:")
    footer(s)

    s = slide(d, 9, "Life returns." + N + " Photos: yellow-caterpillar-on-stem.jpg, caterpillar-on-leaf.jpg, still-young-plants-buckets-29s.png.")
    main(s, B + "yellow-caterpillar-on-stem.jpg", 0.35, 0.4)
    col2(s, B + "caterpillar-on-leaf.jpg", B + "still-young-plants-buckets-29s.png", (0.5, 0.35), (0.5, 0.55))
    words(s, "“I can already see signs of healthier soil in the field, including more earthworms, mushrooms, and active "
             "decomposition.”", voice=True)
    footer(s)

    s = slide(d, 10, "Passing it on." + N + " Photos: group-of-five-farmers.jpeg, two-men-at-shade-house.jpeg [VERIFY who, where, when].")
    main(s, B + "group-of-five-farmers.jpeg", 0.42, 0.5)
    col1(s, B + "two-men-at-shade-house.jpeg", 0.5, 0.35)
    words(s, "“Your training changed the way I see agriculture. I now believe that agriculture is biology, not chemical "
             "fertilizer.”", voice=True)
    footer(s)

    s = slide(d, 11, "The dream and the call to action (YouTube)." + N + " Photo: young-coffee-plant.jpg.")
    main(s, B + "young-coffee-plant.jpg")
    quote_mark(s)
    logo(s, COL_X + 20, PHOTO_Y + PHOTO_H - 290, 240, white=False)
    words(s, "“My dream is to restore degraded land in the Fouta Djallon and help my community learn living-soil practices.”",
          voice=True)
    footer(s, "Subscribe on YouTube: @boubacartidianediallo")

    save(d, "01-10-2026-thu-ig-soil-regenerators-boubacar-soil.pptx")


if __name__ == "__main__":
    build()
