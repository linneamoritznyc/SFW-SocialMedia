"""World Teachers' Day (Monday 5 October 2026): every mentor on the team, one card each.

Roster: Linnea's mentor list, 5 Oct 2026 (14 people). About lines are drafted from the team notes in this repo
(SFW All-Team Meeting, 17 Sep 2026: drafts/2026-10-06-before-soil-food-web.md, drafts/2026-10-01-favourite-microbe-series.md),
the mentor bios collected 30 Sep 2026 and content/community.json. Allison checks and fills them in
(templates/posts/october/05-10-2026-mon-ig-teachers-day-mentors.md). Mentors with no about line yet show name and
role only, so no placeholder text reaches the image.

Paper-cut field notes style: brand green and cream, cut-paper flowers and leaves (no creatures), nothing tilted.
Photo where we have one, otherwise a small flower arrangement.

python3 templates/posts/october/build_teachers_day_v2.py  ->  05-10-2026-mon-ig-teachers-day-v2.pptx
"""
import os
from PIL import Image, ImageFont
from pptx.util import Emu
from lib import Deck, rect, text, crop, ROOT, PX, HEAD, BODY
from build_oct_01_10 import logo, save

GREEN, CREAM, SAGE, PANEL, BLACK = "31662F", "F3F1EA", "B1BCB1", "E4E6DC", "1A1A1A"   # docs/brand-colors.md
GARA = "EB Garamond"
CUT = "assets/collage/cutouts/"
MENT = "assets/mentors-teachers-day/"
W, H, M = 1080, 1350, 80
F = {k: CUT + v for k, v in dict(tithonia="tithonia-flower-1.png", dried="dried-flowers-1-1.png",
                                 tansy="dried-flowers-2-1.png", sprig="leaf-sprig-paper-1.png",
                                 seedling="happy-seedling-1.png", dandelion="dandelion-short-root-1.png",
                                 frame="leaf-frame-1.png").items()}

# first, last, role, about (draft, None = Allison to add), advice (None = none on file), photo, focus, flowers
MENTORS = [
    ("Tommy", "Tepper", "Director of Education & Mentor", None, None, None, None, ("tithonia", "sprig")),
    ("Loida", "Vasquez", "Advanced Programs Lead & Mentor", None, None, "assets/photo/loida-teaching-3.jpg", (0.3, 0.25), ("dried",)),
    ("Carla", "Portugal", "Science Lead & Mentor",
     "PhD in Environmental Sciences. 20 years of environmental and farm consulting. Mentoring since 2019.",
     "Bare soil erodes. Living roots and cover hold the aggregates together.", MENT + "carla-portugal.jpg", (0.5, 0.25), ("tansy",)),
    ("Wesley", "Sanders", "Advanced Programs Mentor",
     "A farmer and agricultural journalist, then our first lab tech. Runs Foothill Biological Soil Health Services in "
     "California's Sierra Nevada foothills.",
     "Compost can look finished and still lack the biology you need. Check it under the microscope before you apply it.",
     MENT + "wes-sander.jpg", (0.5, 0.35), ("sprig",)),
    ("Casey", "Williams", "Advanced Programs Mentor",
     "Around ten years of farming and gardening. A long-time student, now teaching. Favourite microbe: rotifers.",
     None, None, None, ("seedling", "sprig")),
    ("Brian", "Daubenspeck", "Advanced Programs Mentor",
     "Manages orchards in California, after years in New Mexico. Favourite microbes: the small underdogs.",
     None, None, None, ("dried", "tithonia")),
    ("Isadora", "Schmidt", "Advanced Programs Mentor",
     "From Florianópolis, Brazil, now living in Spain. Favourite: lichens, though she loves all her microbes equally.",
     None, None, None, ("tansy", "seedling")),
    ("Aysen", "Ustunay", "Advanced Programs Mentor",
     "Favourite microbe: flagellates. She loves watching them as babies, running around.",
     None, None, None, ("dandelion", "sprig")),
    ("Dora", "Tkalec", "Advanced Programs Mentor", None, None, None, None, ("tithonia", "tansy")),
    ("Gerald", "Ramirez", "Advanced Programs Mentor",
     "An agronomist from the University of Costa Rica, who went looking for Dr. Elaine's teaching. Now he teaches it, "
     "from compost extracts to teas.",
     "Extracts pull organisms off the compost into solution, so you can apply biology across a whole field.",
     MENT + "gerald-ramirez.jpg", (0.5, 0.3), ("dried",)),
    ("Elena", "Kalli", "Advanced Programs Admin", None, None, None, None, ("seedling", "dried")),
    ("Ib", "Borup Pederson", "Advanced Programs Mentor", None, None, None, None, ("sprig", "dandelion")),
    ("Nick", "Padwick", "Advanced Programs Mentor",
     "Manages Ken Hill Estate, home of Wild Ken Hill, in West Norfolk, England. Farmers Weekly Farmer of the Year, 2009.",
     "I make 750 tons of compost a year. Biology works at any scale.", MENT + "Nick.png", (0.5, 0.06), ("tansy",)),
    ("Delvin", "Solkinson", "Permaculture Lead Teacher", None, None, None, None, ("tithonia", "dried", "sprig")),
]

_F = {}
FONTFILE = {(HEAD, True): "~/.fonts/Montserrat-Bold.ttf", (BODY, False): "~/.fonts/SourceSans3-Regular.ttf",
            (GARA, False): "~/.fonts/EBGaramond-Italic.ttf"}


def lines(t, pt, w, font, bold=False):
    k = (font, bold, pt)
    if k not in _F: _F[k] = ImageFont.truetype(os.path.expanduser(FONTFILE[(font, bold)]), int(pt * 1.3333))
    n, cur = 1, ""
    for word in t.split(" "):
        trial = (cur + " " + word).strip()
        if cur and _F[k].getlength(trial) > w * 0.97: n += 1; cur = word
        else: cur = trial
    return n


def para(s, x, y, w, t, sizes, color, font, bold=False, italic=False, sp=1.15, maxh=9999):
    for pt in sizes:
        h = lines(t, pt, w, font, bold) * pt * 1.3333 * 1.17 * sp
        if h <= maxh: break
    text(s, x, y, w, h + 10, t, pt, color, font, bold, italic=italic, spacing=sp)
    return y + h


def piece(s, rel, x, y, w, h):
    """Cut-paper piece fitted inside a box, straight."""
    im = Image.open(os.path.join(ROOT, rel)).convert("RGBA"); im = im.crop(im.split()[3].point(lambda v: 255 if v > 20 else 0).getbbox())
    out = os.path.join(ROOT, "renders", ".tmp", "td-" + os.path.basename(rel)); os.makedirs(os.path.dirname(out), exist_ok=True)
    im.thumbnail((800, 800)); im.save(out)
    r = min(w / im.width, h / im.height); pw, ph = im.width * r, im.height * r
    s.shapes.add_picture(out, Emu(int((x + (w - pw) / 2) * PX)), Emu(int((y + (h - ph) / 2) * PX)), Emu(int(pw * PX)), Emu(int(ph * PX)))


def card(d, first, last, role, about, advice, photo, focus, flowers):
    name = f"{'Dr. ' if last == 'Portugal' else ''}{first} {last}"
    s = d.slide(CREAM, f"Mentor card: {name}, {role}. About line: draft from the team notes, confirm with Allison and "
                       f"{first}. " + (f"Advice line from the mentor bios, 30 Sep 2026: confirm with {first}. " if advice else "")
                       + (f"Photo: {photo} (written consent needed). " if photo else "No photo yet: flower arrangement.")
                       + ("" if about else " [ABOUT LINE NEEDED from Allison; the card shows name and role until then]"),
                counter=False)
    top, ph = M, (560 if (about or advice or photo) else 820)
    if photo:
        s.shapes.add_picture(crop(photo, 920, ph, *focus), Emu(M * PX), Emu(top * PX), Emu(920 * PX), Emu(ph * PX))
        piece(s, F[flowers[0]], W - M - 150, top + ph - 120, 170, 230)          # a flower tucked on the photo corner
    else:
        rect(s, M, top, 920, ph, PANEL)
        n = len(flowers); slot = 920 / n
        for i, f in enumerate(flowers):
            piece(s, F[f], M + i * slot + 20, top + 40, slot - 40, ph - 80)
    y = top + ph + 34
    y = para(s, M, y, 920, name, (50, 46, 42), GREEN, HEAD, True, sp=1.0) + 6
    y = para(s, M, y, 920, role, (22,), BLACK, BODY) + 26
    room = H - 120 - y
    if advice:
        y = para(s, M, y, 920, f"“{advice}”", (38, 34, 30, 28), BLACK, GARA, italic=True, maxh=room * 0.55) + 20
        room = H - 120 - y
    if about:
        para(s, M, y, 920, about, (26, 24, 22, 20), GREEN if advice else BLACK, BODY, maxh=room)
    text(s, W - M - 300, H - 70, 300, 30, "soilfoodweb.com", 18, GREEN, HEAD, True, align="r")


def build():
    d = Deck(name="05-10-2026-mon-ig-teachers-day-v2")

    s = d.slide(GREEN, "Cover. World Teachers' Day, 5 October. Cut-paper flowers from the collage library.", counter=False)
    text(s, M, 120, W - 2 * M, 40, "WORLD TEACHERS' DAY · 5 OCTOBER", 20, SAGE, HEAD, True, track=3)
    text(s, M, 180, W - 2 * M, 360, "Happy World Teachers' Day.", 76, CREAM, HEAD, True, spacing=1.0)
    text(s, M, 540, W - 2 * M, 160, "Meet the mentors who teach our students to see soil.", 36, CREAM, GARA,
         italic=True, spacing=1.15)
    for i, f in enumerate(["tithonia", "dried", "tansy", "sprig", "seedling"]):
        piece(s, F[f], M + i * 184, 780, 170, 380)
    logo(s, W - M - 110, H - 70 - 110 * 668 / 743, 110, white=True)

    for m in MENTORS:
        card(d, *m)

    s = d.slide(GREEN, "Closing: every mentor on the roster thanked by name.", counter=False)
    text(s, M, 140, W - 2 * M, 260, "Thank you to every mentor who teaches our students to see soil.", 52, CREAM, HEAD,
         True, spacing=1.08)
    names = " · ".join(f"{m[0]} {m[1]}" for m in MENTORS)
    text(s, M, 470, W - 2 * M, 380, names, 30, CREAM, GARA, italic=True, spacing=1.35)
    for i, f in enumerate(["sprig", "tansy", "tithonia", "dried"]):
        piece(s, F[f], M + i * 190, 930, 170, 300)
    logo(s, W - M - 110, H - 70 - 110 * 668 / 743, 110, white=True)
    save(d, "05-10-2026-mon-ig-teachers-day-v2.pptx")


if __name__ == "__main__":
    build()
