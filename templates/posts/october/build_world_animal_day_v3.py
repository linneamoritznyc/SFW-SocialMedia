"""World Animal Day carousel, version 3: research style (Sunday 4 October 2026, Instagram, 10 slides, 1080 x 1350).

A field-guide / journal look on cream: a running header and footer rule on every slide, specimen tiles with the animal's
name, a colour-coded legend (green: goes in, gold: hot pile only, red: stays out), numbered footnote citations, a
stacked-bar chart of the feedstock recipe and a drawn pile cross-section. Brand palette (docs/brand-colors.md).
The copy is the brief's, word for word, split over the slides. Photos: the first version's (assets/photo/animals/).

python3 build_world_animal_day_v3.py  ->  04-10-2026-sun-ig-world-animal-day-v3.pptx
Every element is a separate editable object.
"""
import os
from PIL import Image
from pptx.util import Emu, Pt
from lib import Deck, rect, text, crop, poly, est_lines, ROOT, PX, HEAD, BODY, rgb
from build_oct_01_10 import logo, save
from build_world_animal_day_v2 import SOURCES, CRED, cred_line, cred_note

W, H, M = 1080, 1350, 80
CW = W - 2 * M
GREEN, CREAM, BROWN, TAN, SAGE, GOLD, RED, INK = "31662F", "F3F1EA", "4C3634", "C09D7F", "B1BCB1", "D39C48", "B23A2E", "333130"
AN = "assets/photo/animals/"
ICON = "assets/icons/icon-{}-{}.png"
CAT = {"in": (GREEN, "GOES IN"), "hot": (GOLD, "HOT PILE ONLY"), "out": (RED, "STAYS OUT")}
FOOT_Y = H - 128                      # footer rule


# ---------------------------------------------------------------- page furniture
def page(d, note):
    s = d.slide(CREAM, note, counter=False)
    text(s, M, 56, 600, 24, "WORLD ANIMAL DAY, 4 OCTOBER 2026", 14, BROWN, HEAD, True, track=2)
    text(s, W - M - 300, 56, 300, 24, "soilfoodweb.com", 14, BROWN, HEAD, True, align="r")
    rule(s, 92)
    return s


def rule(s, y, x=M, w=CW, alpha=35, h=2, color=BROWN):
    rect(s, x, y, w, h, color, alpha=alpha)


def footnotes(s, notes):
    """Numbered sources under a rule at the foot of the page, colour logo bottom right."""
    rule(s, FOOT_Y)
    y = FOOT_Y + 16
    for i, n in enumerate(notes, 1):
        rich(s, y, [(str(i), "sup"), ("  " + n, "")], pt=14, color=INK, alpha=75, w=CW - 120, sp=1.1)
        y += 24
    logo(s, W - M - 76, FOOT_Y + 18, 76, white=False)


def rich(s, y, runs, pt=30, color=INK, x=M, w=CW, sp=1.3, alpha=None, h=None):
    """One paragraph of runs: (text, style) with style '' (Source Sans 3), 'b' (Montserrat Bold) or 'sup' (superscript)."""
    full = "".join(t for t, _ in runs)
    h = h or est_lines(full, pt, w * 0.95) * pt * 1.3333 * sp + 12
    tb = s.shapes.add_textbox(0, 0, Emu(1), Emu(1)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.line_spacing = sp
    for t, st in runs:
        r = p.add_run(); r.text = t; f = r.font
        f.size = Pt(pt); f.name = HEAD if st == "b" else BODY; f.bold = st == "b"; f.color.rgb = rgb(color)
        if st == "sup": r._r.get_or_add_rPr().set("baseline", "30000")
        if alpha is not None:
            from lib import _alpha, A
            _alpha(r._r.get_or_add_rPr().find("{%s}solidFill" % A), alpha)
    tb.left, tb.top, tb.width, tb.height = Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX))
    return y + h


def head(s, y, t, pt=60, color=GREEN, w=CW, lines=None):
    n = lines or est_lines(t, pt, w * 0.92, True)
    h = n * pt * 1.3333 * 1.12 + 6
    text(s, M, y, w, h, t, pt, color, HEAD, True, spacing=1.05)
    return y + h


def marker(s, y, cat, x=M):
    col, lab = CAT[cat]
    rect(s, x, y + 3, 22, 22, col)
    text(s, x + 36, y, 600, 28, lab, 18, INK, HEAD, True, track=2, anchor="m")


def pic(s, key, x, y, w, h, fx=0.5, fy=0.5):
    p = crop(AN + key + ".jpg", int(w), int(h), fx, fy)
    s.shapes.add_picture(p, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))
    rect(s, x, y, w, h, None, line=BROWN, lw=1)          # hairline figure frame


def caption(s, y, t, x=M, w=CW):
    text(s, x, y, w, 24, t, 14, INK, BODY, italic=True, alpha=75)


def icon(s, name, colour, x, y, h):
    p = os.path.join(ROOT, ICON.format(name, colour))
    iw, ih = Image.open(p).size; w = h * iw / ih
    return s.shapes.add_picture(p, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX))), w


FOCUS = {"rabbit": (0.5, 0.45), "goat": (0.55, 0.45), "sheep": (0.35, 0.5), "alpaca": (0.5, 0.3), "guinea-pig": (0.45, 0.5),
         "horse": (0.5, 0.55), "cow": (0.5, 0.35), "chicken": (0.5, 0.5), "pig": (0.7, 0.45), "dog": (0.5, 0.4), "cat": (0.62, 0.6)}
NAME = {"guinea-pig": "Guinea pig"}


def specimen(s, key, x, y, w, ph, cat):
    """Specimen tile: photo, then a name strip with the category colour square."""
    pic(s, key, x, y, w, ph, *FOCUS[key])
    rect(s, x, y + ph, w, 34, CREAM)
    rect(s, x + 2, y + ph + 11, 12, 12, CAT[cat][0])
    text(s, x + 22, y + ph + 2, w - 24, 30, NAME.get(key, key.capitalize()), 16, INK, HEAD, True, anchor="m")


# ---------------------------------------------------------------- slides
def build():
    d = Deck(name="04-10-2026-sun-ig-world-animal-day-v3")
    herb = ["rabbit", "goat", "sheep", "alpaca", "guinea-pig", "horse", "cow", "chicken"]
    allk = herb + ["pig", "dog", "cat"]

    # 1 cover: classification chart of eleven animals and the legend
    s = page(d, "Cover, research style. " + SOURCES + " " + cred_note(*allk))
    y = head(s, 124, "Whose manure belongs in your compost?", 54, lines=2) + 20
    text(s, M, y, CW, 44, "Happy World Animal Day, October 4.", 30, INK, BODY)
    g = 10; tw = (CW - 3 * g) / 4; ph = tw - 34; gy = y + 80
    cats = {k: "in" for k in herb}; cats.update(pig="hot", dog="out", cat="out")
    for i, k in enumerate(allk):
        specimen(s, k, M + (i % 4) * (tw + g), gy + (i // 4) * (tw + g), tw, ph, cats[k])
    lx, ly = M + 3 * (tw + g), gy + 2 * (tw + g)                 # legend in the twelfth cell
    rect(s, lx, ly, tw, tw, None, line=BROWN, lw=1)
    for j, c in enumerate(("in", "hot", "out")):
        col, lab = CAT[c]
        rect(s, lx + 18, ly + 34 + j * 62, 18, 18, col)
        text(s, lx + 46, ly + 26 + j * 62, tw - 56, 60, lab, 13, INK, HEAD, True, track=1, spacing=1.0)
    text(s, M, FOOT_Y + 16, CW - 120, 48, cred_line(*allk), 12, INK, BODY, italic=True, alpha=70, spacing=1.1)
    rule(s, FOOT_Y)
    logo(s, W - M - 76, FOOT_Y + 18, 76, white=False)

    # 2 manure: a high-nitrogen feedstock
    s = page(d, "Feedstocks, text, footnote 1.")
    icon(s, "manure-fork", "brown", W - M - 130, 140, 130)
    y = head(s, 150, "Manure: a high-nitrogen feedstock", 64, w=CW - 170, lines=3) + 30
    rule(s, y, w=120, alpha=100, h=4, color=GREEN)
    rich(s, y + 40, [("A Soil Food Web thermophilic compost recipe combines three feedstocks: high-nitrogen material such as "
                      "manure (about 10% in summer), green plant material (about 30%) and woody material (about 60%). Bedding, "
                      "such as straw or wood shavings, counts toward the woody share.", ""), ("1", "sup")], pt=34)
    footnotes(s, ["Soil Food Web School, BioComplete Compost, Lecture 4"])

    # 3 the recipe as a 100% stacked bar with an axis and a legend
    s = page(d, "Feedstocks, chart: one 100% stacked bar (10, 30, 60) with a 0 to 100% axis and a legend. Icons: assets/icons/.")
    parts = [("manure-fork", "High nitrogen, 10%", BROWN, 0.10), ("leaf", "Green, 30%", GREEN, 0.30), ("wood-log", "Woody, 60%", TAN, 0.60)]
    bx, by, bh = M, 330, 150
    x = bx
    for ico, lab, col, sh in parts:
        rect(s, x, by, CW * sh - 4, bh, col); x += CW * sh
    for t in range(0, 101, 25):                                    # axis
        tx = bx + CW * t / 100
        rect(s, tx - 1, by + bh + 12, 2, 14, INK, alpha=60)
        text(s, tx - 40, by + bh + 32, 80, 24, f"{t}%", 16, INK, BODY, align="c", alpha=75)
    rule(s, by + bh + 12, alpha=60)
    for j, (ico, lab, col, sh) in enumerate(parts):                # legend
        ly = 640 + j * 150
        rect(s, M, ly + 22, 34, 34, col)
        icon(s, ico, "brown", M + 70, ly, 80)
        text(s, M + 240, ly + 6, CW - 240, 70, lab, 40, INK, HEAD, True, anchor="m")
        if j < 2: rule(s, ly + 118, alpha=20)
    text(s, M, 150, CW, 60, "High nitrogen, 10% · Green, 30% · Woody, 60%", 24, GREEN, HEAD, True, track=1)
    footnotes(s, ["Soil Food Web School, BioComplete Compost, Lecture 4"])

    # 4 goes in: plant-eaters and poultry
    s = page(d, "Plant-eaters and poultry. " + cred_note(*herb))
    g = 10; tw = (CW - 3 * g) / 4; ph = tw - 34
    for i, k in enumerate(herb):
        specimen(s, k, M + (i % 4) * (tw + g), 130 + (i // 4) * (tw + g), tw, ph, "in")
    y = 130 + 2 * tw + g + 10
    text(s, M, y, CW, 40, cred_line(*herb), 12, INK, BODY, italic=True, alpha=70, spacing=1.1)
    marker(s, y + 62, "in")
    y = head(s, y + 108, "Plant-eaters and poultry", 54, lines=1) + 6
    rich(s, y, [("Rabbit, goat, sheep, alpaca, guinea pig, horse, cow and chicken. Their manure is a great high-nitrogen "
                 "ingredient. Use it fresh, or dry it and store it until you build the pile.", ""), ("1", "sup")])
    footnotes(s, ["Soil Food Web School, Compost Manual"])

    # 5 to 7: one animal per page, photo as a figure
    def animal(key, cat, title, runs, notes, note):
        s = page(d, note + " " + cred_note(key))
        pic(s, key, M, 130, CW, 520, *FOCUS[key])
        caption(s, 662, cred_line(key))
        marker(s, 712, cat)
        y = 760
        if title: y = head(s, y, title, 60, lines=1) + 8
        rich(s, y, runs)
        footnotes(s, notes)
    animal("pig", "hot", "Pig", [("Pig roundworm can infect people.", ""), ("1", "sup"), (" Pig manure goes into a pile that "
           "completes the thermophilic requirements, and raw pig manure stays away from produce.", "")],
           ["Miller et al. 2015, Emerging Infectious Diseases"], "Pig.")
    animal("dog", "out", "Dog and cat", [("Dog feces can carry roundworm (Toxocara) eggs. If swallowed, the larvae can migrate "
           "to the liver, lungs and eyes.", ""), ("1", "sup")], ["USDA NRCS 2005; CDC 2025"], "Dog and cat, part 1.")
    animal("cat", "out", None, [("Cat feces can carry Toxoplasma gondii, which people can also pick up from contaminated soil.", ""),
           ("2", "sup"), (" USDA advises against using dog waste compost on food crops, even after hot composting, and against "
           "composting cat waste or litter at all.", ""), ("1", "sup")], ["USDA NRCS 2005", "CDC 2025"], "Dog and cat, part 2.")

    # 8 slow or hot: two columns, a figure under each
    s = page(d, "Slow or hot, comparison in two columns. " + cred_note("compost-2", "compost-4"))
    head(s, 130, "Slow or hot, both work.", 60, lines=2)
    cw2 = (CW - 40) / 2
    for i, (lead, rest, key, fy) in enumerate([("Cold or static pile:", " fine for plant-eater and poultry manure, on a longer "
                                                "timeline.", "compost-2", 0.75),
                                               ("Thermophilic pile:", " above 55°C (131°F) for 3 days at the center.", "compost-4", 0.6)]):
        x = M + i * (cw2 + 40)
        rect(s, x, 350, cw2, 6, (SAGE, GOLD)[i])
        rich(s, 380, [(lead, "b"), (rest, ""), ("1", "sup")], pt=30, x=x, w=cw2, h=240)
        pic(s, key, x, 670, cw2, 450, 0.5, fy)
        caption(s, 1130, cred_line(key), x=x, w=cw2)
    footnotes(s, ["Soil Food Web School, Compost Manual"])

    # 9 turn it: pile cross-section with the hot center
    s = page(d, "Thermophilic pile, diagram: a pile cross-section drawn with shapes (tan pile, gold hot center), "
                "arrows show turning the outside into the center. Thermometer icon: assets/icons/.")
    rich(s, 130, [("Thermophilic pile:", "b"), (" above 55°C (131°F) for 3 days at the center. Turn it so every part passes "
                  "through the hot center.", ""), ("1", "sup")], pt=34, w=CW)
    import math
    cx, base, rx, ry = W / 2, 1080, 440, 470
    pts = [(cx + rx * math.cos(math.pi * t / 60), base - ry * math.sin(math.pi * t / 60)) for t in range(61)]
    poly(s, pts, TAN)
    pts2 = [(cx + 300 * math.cos(math.pi * t / 60), base - 320 * math.sin(math.pi * t / 60)) for t in range(61)]
    poly(s, pts2, "D8C3AE")
    core = [(cx + 170 * math.cos(2 * math.pi * t / 60), base - 170 + 120 * math.sin(2 * math.pi * t / 60)) for t in range(60)]
    poly(s, core, GOLD)
    text(s, cx - 170, base - 230, 340, 50, "above 55°C (131°F)", 26, BROWN, HEAD, True, align="c", anchor="m")
    text(s, cx - 170, base - 180, 340, 40, "3 days", 22, BROWN, BODY, align="c", anchor="m")
    rect(s, M, base, CW, 4, BROWN)
    for side in (-1, 1):                                           # turning arrows: outside layer into the center
        x0 = cx + side * 390; y0 = base - 120
        x1 = cx + side * 200; y1 = base - 170
        L = math.hypot(x1 - x0, y1 - y0); ux, uy = (x1 - x0) / L, (y1 - y0) / L; nx, ny = -uy, ux
        poly(s, [(x0 + nx * 5, y0 + ny * 5), (x1 - ux * 26 + nx * 5, y1 - uy * 26 + ny * 5), (x1 - ux * 26 - nx * 5, y1 - uy * 26 - ny * 5),
                 (x0 - nx * 5, y0 - ny * 5)], BROWN)
        poly(s, [(x1, y1), (x1 - ux * 30 + nx * 16, y1 - uy * 30 + ny * 16), (x1 - ux * 30 - nx * 16, y1 - uy * 30 - ny * 16)], BROWN)
    icon(s, "compost-thermometer", "brown", cx - 40, base - 560, 120)
    footnotes(s, ["Soil Food Web School, Compost Manual"])

    # 10 the course
    s = page(d, "Closing. Course: https://school.soilfoodweb.com/bundles/advanced-biocomplete-compost-production")
    head(s, 170, "Advanced BioComplete Compost Production", 64, lines=4)
    rule(s, 640, w=120, alpha=100, h=4, color=GREEN)
    text(s, M, 680, CW, 200, "Building and assessing three compost piles that meet biological minimums. At the Soil Food Web "
                             "School. Link in bio.", 32, INK, BODY, spacing=1.3)       # course line from docs/copy-deck-v2.md
    lw = 220
    logo(s, (W - lw) / 2, 940, lw, white=False)
    rule(s, FOOT_Y)
    save(d, "04-10-2026-sun-ig-world-animal-day-v3.pptx")


if __name__ == "__main__":
    build()
