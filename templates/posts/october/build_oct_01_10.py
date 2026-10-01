"""October 1 to 10, 2026: every visual except Boubacar (skipped) and World Animal Day (built by build_world_animal_day.py).

python3 build_oct_01_10.py [post-number ...]   ->   templates/posts/october/DD-10-2026-*.pptx
Colours, fonts and templates from lib.py (the October spec). Every element is a separate editable object.
No slide numbers on any post.
"""
import os, sys
from PIL import Image
from pptx.util import Emu
from lib import (Deck, Tilt, _place, rect, rrect, oval, text, gradient, crop, fit, est_lines, ROOT, PX,
                 DEEP, GREEN, CREAM, GLOW, LIGHT, INK, FAINT, STRIPE, HEAD, BODY)

HERE = os.path.dirname(os.path.abspath(__file__))
GRAIN = "assets/collage/paper-grain-cream.jpg"          # 1080 x 1350, made by build_world_animal_day.py
CUT = "assets/collage/cutouts/"
MENT = "assets/mentors-teachers-day/"
LOGO_COLOR = "assets/logo/foundation-logo-color.png"     # 743 x 668
LOGO_WHITE = "assets/logo/foundation-logo-white.png"
RED = "B23A2E"
TMP = os.path.join(HERE, "..", "..", "..", "renders", ".tmp")


# ---------------------------------------------------------------- shared pieces
def pic(s, rel, x, y, w, h=None, deg=0):
    p = os.path.join(ROOT, rel)
    if h is None:
        iw, ih = Image.open(p).size; h = w * ih / iw
    sh = s.shapes.add_picture(p, 0, 0, Emu(1), Emu(1)); _place(sh, x, y, w, h); sh.rotation = deg
    return sh


def trimmed(rel):
    """Cutout cropped to its visible pixels, so the box hugs the piece."""
    os.makedirs(TMP, exist_ok=True)
    out = os.path.join(TMP, "trim-" + os.path.basename(rel))
    im = Image.open(os.path.join(ROOT, rel)).convert("RGBA")
    im.crop(im.split()[3].point(lambda v: 255 if v > 20 else 0).getbbox()).save(out)
    return out


def piece(s, rel, cx, cy, w, deg=0):
    """Collage cutout centred at (cx, cy), w px wide. Silently skipped if the cutout was never generated."""
    if not os.path.exists(os.path.join(ROOT, rel)):
        MISSING.add(rel); return None
    p = trimmed(rel); iw, ih = Image.open(p).size; h = w * ih / iw
    sh = s.shapes.add_picture(p, 0, 0, Emu(1), Emu(1)); _place(sh, cx - w / 2, cy - h / 2, w, h); sh.rotation = deg
    return sh


def fitpiece(s, rel, x, y, w, h, deg=0):
    """Cutout scaled to fit inside the box (x, y, w, h), centred."""
    if not os.path.exists(os.path.join(ROOT, rel)):
        MISSING.add(rel); return None
    iw, ih = Image.open(trimmed(rel)).size; r = min(w / iw, h / ih)
    return piece(s, rel, x + w / 2, y + h / 2, iw * r, deg)


def logo(s, x, y, w=90, white=True):
    return pic(s, LOGO_WHITE if white else LOGO_COLOR, x, y, w, w * 668 / 743)


def logo_br(s, white=True, w=90, W=1080, H=1350, m=60):
    return logo(s, W - m - w, H - m - w * 668 / 743, w, white)


def logo_box(s, x, y, w, h, radius=0):
    """Photo box with no photo yet: stripe paper with the colour logo in the middle."""
    (rrect(s, x, y, w, h, CREAM, radius=radius, pattern=STRIPE) if radius else rect(s, x, y, w, h, CREAM, pattern=STRIPE))
    lw = min(w, h) * 0.55
    logo(s, x + (w - lw) / 2, y + (h - lw * 668 / 743) / 2, lw, white=False)


def photo_box(s, rel, x, y, w, h, fx=0.5, fy=0.5, radius=0):
    """Real photo cropped to the box, or the logo box when rel is None."""
    if rel is None:
        return logo_box(s, x, y, w, h, radius)
    p = crop(rel, int(w), int(h), fx, fy)
    sh = s.shapes.add_picture(p, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))
    if radius:
        from pptx.enum.shapes import MSO_SHAPE
        from lxml import etree
        A = "http://schemas.openxmlformats.org/drawingml/2006/main"
        sh.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        g = sh._element.spPr.find("{%s}prstGeom" % A); av = g.find("{%s}avLst" % A)
        if av is None: av = etree.SubElement(g, "{%s}avLst" % A)
        etree.SubElement(av, "{%s}gd" % A, name="adj", fmla="val %d" % int(min(50000, radius / min(w, h) * 100000)))
    return sh


def cream(d, note, w=1080, h=1350):
    s = d.slide(CREAM, note, counter=False)
    s.shapes.add_picture(os.path.join(ROOT, GRAIN), 0, 0, Emu(w * PX), Emu(h * PX))
    return s


def dark(d, note, bg=DEEP):
    return d.slide(bg, note, counter=False)


def lh(pt, spacing=1.2):
    """Rendered line height in px (LibreOffice and PowerPoint set about 1.2 x the size per line)."""
    return pt * 1.3333 * 1.2 * spacing


from PIL import ImageFont
FONTS = {(HEAD, True): "Montserrat-Bold.ttf", (HEAD, False): "Montserrat-Regular.ttf",
         (BODY, True): "SourceSans3-Bold.ttf", (BODY, False): "SourceSans3-Regular.ttf",
         ("EB Garamond", False): "EBGaramond-Italic.ttf", ("EB Garamond", True): "EBGaramond-Italic.ttf"}


def text_width(t, pt, font=HEAD, bold=True):
    f = ImageFont.truetype(os.path.expanduser("~/.fonts/" + FONTS[(font, bold)]), int(pt * 1.3333))
    return f.getlength(t)


def n_lines(t, pt, w, font=HEAD, bold=True):
    """Lines after word wrap, measured with the real brand fonts (installed in ~/.fonts)."""
    n = 0
    for para in t.split("\n"):
        n += 1; line = ""
        for word in para.split(" "):
            trial = (line + " " + word).strip()
            if line and text_width(trial, pt, font, bold) > w: n += 1; line = word
            else: line = trial
    return n


def block(s, x, y, w, t, pt, color=DEEP, font=HEAD, bold=True, spacing=1.15, align="l", italic=None):
    """Text box sized from the measured number of lines; returns the bottom y. EB Garamond is set in italic."""
    n = n_lines(t, pt, w * 0.97, font, bold)
    h = n * lh(pt, spacing)
    if italic is None: italic = font == "EB Garamond"
    text(s, x, y, w, h + 10, t, pt, color, font, bold, italic=italic, align=align, spacing=spacing)
    return y + h


def centered(s, t, pt, color=DEEP, rule=GREEN, cy=675):
    """Big statement set around the vertical middle, with a short rule above it."""
    top = cy - n_lines(t, pt, 920 * 0.97) * lh(pt, 1.12) / 2
    if rule: rect(s, 80, top - 60, 140, 10, rule)
    return block(s, 80, top, 920, t, pt, color, HEAD, True, 1.12)


def head_body(s, head, body, y=110, hpt=66, bpt=38, w=920, x=80):
    """Punchy headline in Montserrat, then a body line in Source Sans. Returns the bottom y."""
    y = block(s, x, y, w, head, hpt, DEEP, HEAD, True, 1.08) + 26
    return block(s, x, y, w, body, bpt, INK, BODY, False, 1.25)


def source_line(s, t, dark_bg=False, W=1080, H=1350, right_gap=200):
    text(s, 80, H - 95, W - 80 - right_gap, 40, t, 18, GLOW if dark_bg else FAINT, BODY, alpha=85 if dark_bg else None)


def unsplash_credit(fname):
    for line in open(os.path.join(ROOT, "assets/photo/unsplash-credits.txt")):
        f = line.rstrip("\n").split("\t")
        if f[0] == fname:
            return f"{f[1]} ({f[2].replace('Profile: ', '')})."
    return "[CREDIT MISSING]"


MISSING = set()
BUILT = []


def save(d, fname):
    path = os.path.join(HERE, fname)
    d.finish(path, counters=False)
    BUILT.append(fname)
    print("wrote", fname)


# ================================================================ 1. Coffee myth (Thu 1 Oct), template 2
def coffee():
    DRAFT = " Draft livelier copy (not in use, awaiting approval): "
    """Cute science scrapbook: taped prints, drawn stickers, cream paper."""
    import scrapbook as sb
    d = Deck(name="01-10-2026-thu-ig-coffee-myth")
    src = ("Sources: Hardgrove and Livesley, Urban Forestry and Urban Greening, 2016, https://doi.org/10.1016/j.ufug.2016.02.015; "
           "Summers et al., Journal of Bacteriology, 2012, https://doi.org/10.1128/JB.06637-11; "
           "US EPA, 40 CFR Part 503, https://www.ecfr.gov/current/title-40/chapter-I/subchapter-O/part-503.")
    cr = lambda f: "Photo: " + unsplash_credit(f).split(" (")[0].replace("Photo by ", "")
    stickers = " Stickers (lemon, cup, beans, seedlings, bacteria, pH strip, thermometer) are drawn PowerPoint shapes, grouped and editable."

    # 1 myth
    s = cream(d, "Myth slide. " + src + " Photos: " + unsplash_credit("coffee-cup-leaves-top-view.jpg") + " "
              + unsplash_credit("coffee-grounds-close-up.jpg") + stickers)
    x, y, fw, fh = sb.polaroid(s, "assets/photo/coffee-cup-leaves-top-view.jpg", 70, 90, 470, -5, cr("coffee-cup-leaves-top-view.jpg"))
    sb.tape(s, x + fw / 2, y + 8, 160, 46, -6)
    x2, y2, fw2, fh2 = sb.polaroid(s, "assets/photo/coffee-grounds-close-up.jpg", 560, 150, 410, 4, cr("coffee-grounds-close-up.jpg"))
    sb.tape(s, x2 + fw2 - 50, y2 + 10, 140, 44, 32)
    sb.lemon(s, 960, 130, 85, 12)
    sb.beans(s, 520, 690, 170, -10)
    tl = sb.note(s, 70, 770, 940, 430, DEEP, -1.5)
    text(s, 120, 810, 840, 90, "Myth:", 60, GLOW, HEAD, True, anchor="m", tl=tl)
    for i, ln in enumerate(["“Coffee grounds", "acidify your soil.”"]):
        yy = 905 + i * 115
        text(s, 120, yy, 900, 110, ln, 70, CREAM, HEAD, True, anchor="m", spacing=1.0, tl=tl)
        rect(s, 106, yy + 55, text_width(ln, 70) + 28, 16, GLOW, tl=tl)
    s.notes_slide.notes_text_frame.text += DRAFT + "Happy International Coffee Day ☕ Let's spill the beans."
    logo_br(s, False, 110)

    # 2 neutral pH
    s = cream(d, "pH strip: universal indicator colours, 0 to 14. Lemon marks the acid end; the cup sits near neutral, as the slide says." + stickers)
    s.notes_slide.notes_text_frame.text += DRAFT + "Plot twist: your coffee took the acid with it. Brewing pulls most of the acids into your cup. The spent grounds left behind sit close to neutral pH."
    yb = head_body(s, "Brewing extracts most of the acids.", "Spent grounds sit close to neutral pH.", hpt=62, bpt=44)
    sy = yb + 250
    cw = sb.ph_strip(s, 90, sy, 900)
    sb.lemon(s, 90 + 2.5 * cw, sy - 120, 70, -10); sb.pointer(s, 90 + 2.5 * cw, sy - 40)
    sb.cup(s, 90 + 6.5 * cw, sy - 130, 160); sb.pointer(s, 90 + 6.5 * cw, sy - 40)
    sb.beans(s, 250, 1150, 180, 8)
    sb.cutout(s, CUT + "dried-flowers-2-1.png", 940, 1230, 170, 22)

    # 3 caffeine and chlorogenic acids
    s = cream(d, "Photo: " + unsplash_credit("seedlings-sprouting-in-soil.jpg") + stickers)
    s.notes_slide.notes_text_frame.text += DRAFT + "So what's the catch? Caffeine. 😬 Leftover caffeine and chlorogenic acids are allelopathic: plant-speak for chemicals that stop seeds sprouting and roots growing."
    head_body(s, "The real issue: residual caffeine and chlorogenic acids.", "Both are allelopathic. They inhibit germination and root growth.",
              hpt=58, bpt=44)
    x, y, fw, fh = sb.polaroid(s, "assets/photo/seedlings-sprouting-in-soil.jpg", 90, 700, 440, -4, cr("seedlings-sprouting-in-soil.jpg"))
    sb.tape(s, x + 60, y + 8, deg=-30)
    sb.seedling(s, 800, 1080, 300, sad=True)
    sb.beans(s, 690, 1180, 150, 15)
    sb.cup(s, 950, 660, 140, 8, happy=False)

    # 4 fresh grounds in trials
    s = cream(d, "Photo: assets/photo/garden-vegetable-beds.jpg (Foundation library)." + stickers)
    s.notes_slide.notes_text_frame.text += DRAFT + "Sprinkled straight on the garden? In trials, fresh grounds spread on beds actually slowed plant growth. 🥀"
    block(s, 80, 110, 920, "Fresh grounds spread on beds have reduced plant growth in trials.", 58, DEEP, HEAD, True, 1.08)
    x, y, fw, fh = sb.polaroid(s, "assets/photo/garden-vegetable-beds.jpg", 330, 620, 500, 3, "Foundation library")
    sb.tape(s, x + fw - 60, y + 10, deg=30); sb.tape(s, x + 60, y + 10, deg=-30)
    sb.seedling(s, 190, 1010, 240, sad=True, deg=-4)
    sb.beans(s, 200, 1180, 170, -12)
    sb.cup(s, 900, 1200, 140, -8, happy=False)

    # 5 bacteria that eat caffeine
    s = cream(d, "Microscopy: assets/microscopy/sfw-amoeba-still-square.jpg (Foundation). The library has no bacteria-only image; "
                 "this still shows an amoeba, a bacteria feeder. Swap in a bacteria micrograph if Wes has one [IMAGE NEEDED]. "
                 "The drawn bacteria are rods with a flagellum, like Pseudomonas." + stickers)
    x, y, fw, fh = sb.polaroid(s, "assets/microscopy/sfw-amoeba-still-square.jpg", 80, 80, 520, -3, "Soil Food Web Foundation")
    sb.tape(s, x + fw / 2, y + 6, 160, 46, 4)
    sb.bacterium(s, 820, 200, 220, -18); sb.bacterium(s, 880, 440, 190, 12, "D6E3A8"); sb.bacterium(s, 760, 640, 170, -6, "C9DDB6")
    sb.beans(s, 990, 590, 110, 20)
    s.notes_slide.notes_text_frame.text += DRAFT + "Enter the compost crew. 🦠 Composting breaks caffeine down. Some soil bacteria eat it for breakfast, lunch and dinner: Pseudomonas putida CBB5 uses caffeine as its only carbon and nitrogen source."
    head_body(s, "Composting breaks them down.",
              "Some soil bacteria, like Pseudomonas putida CBB5, use caffeine as their only carbon and nitrogen source.", y=790, hpt=62, bpt=42)

    # 6 compost them hot
    s = cream(d, "Photo: assets/photo/hand-of-compost.jpg (Foundation library). Thermometer marks 55°C, from the slide text." + stickers)
    s.notes_slide.notes_text_frame.text += DRAFT + "The fix: compost them first. 🔥 Grounds are a nitrogen-rich green. Mix them with woody browns like leaves or wood chips, and keep the pile hot: 55°C (131°F) or more."
    head_body(s, "Grounds are a green, nitrogen-rich input.", "Mix with woody browns and compost them hot, at 55°C or more.", hpt=62, bpt=44)
    x, y, fw, fh = sb.polaroid(s, "assets/photo/hand-of-compost.jpg", 90, 700, 460, -4, "Foundation library")
    sb.tape(s, x + 70, y + 8, deg=-30)
    my = sb.thermometer(s, 690, 700, 520, 55)
    text(s, 780, my - 32, 260, 64, "55°C", 44, DEEP, HEAD, True, anchor="m")
    sb.cup(s, 920, 1230, 150, 6)

    # 7 sources
    s = cream(d, "Sources slide.")
    tl = sb.note(s, 70, 150, 940, 760, sb.PAPER, -1)
    text(s, 120, 200, 840, 100, "Sources.", 64, DEEP, HEAD, True, anchor="m", tl=tl)
    yy = 330
    for t in ("Hardgrove and Livesley, Urban Forestry and Urban Greening, 2016: https://doi.org/10.1016/j.ufug.2016.02.015",
              "Summers et al., Journal of Bacteriology, 2012: https://doi.org/10.1128/JB.06637-11",
              "US EPA, 40 CFR Part 503: https://www.ecfr.gov/current/title-40/chapter-I/subchapter-O/part-503"):
        n = n_lines(t, 22, 820 * 0.97, BODY, False)
        text(s, 120, yy, 840, n * lh(22, 1.25) + 10, t, 22, INK, BODY, False, spacing=1.25, tl=tl); yy += n * lh(22, 1.25) + 60
    s.notes_slide.notes_text_frame.text += DRAFT + "Save this for your next coffee run. ☕"
    sb.lemon(s, 330, 1160, 80, -12); sb.cup(s, 570, 1170, 170, 4); sb.bacterium(s, 790, 1150, 170, -10)
    sb.cutout(s, CUT + "dried-flowers-2-1.png", 120, 1210, 170, -25)
    logo_br(s, False, 110)
    save(d, "01-10-2026-thu-ig-coffee-myth.pptx")


# ================================================================ 2 and 3. 1000 Farms study (Fri 2 Oct)
STUDY_SRC = ("Sources: Lundgren et al., Environmental Research: Food Systems, 2026, https://doi.org/10.1088/2976-601X/ae8f4e; "
             "Ecdysis Foundation press release, 17 September 2026, https://www.globenewswire.com/news-release/2026/09/17/3363925/0/en/"
             "regenerative-farming-outpaces-tech-based-carbon-capture-solutions-in-first-large-scale-systemic-study.html")
STUDY_LINE = "Lundgren et al. 2026, Environ. Res.: Food Syst."
FIELD = CUT + "field-cross-section-1.png"


def field(s, x, y, w, h, generated=False):
    """Field collage from existing cut-paper pieces: ferns above for the cover crop, the fungal-hyphae cutout below
    for fungal threads and roots. The generated field-cross-section-1.png is the banned soil cross-section style
    (brown layers, grass fringe), so it is only used when asked for explicitly."""
    if generated and os.path.exists(os.path.join(ROOT, FIELD)):
        return fitpiece(s, FIELD, x, y, w, h)
    iw, ih = Image.open(trimmed("assets/collage/cut-brown-brush.png")).size       # torn-paper soil layer
    soil = s.shapes.add_picture(trimmed("assets/collage/cut-brown-brush.png"), 0, 0, Emu(1), Emu(1))
    _place(soil, x + w * 0.04, y + h * 0.4, w * 0.92, h * 0.56)
    piece(s, "assets/collage/fern-green.png", x + w * 0.32, y + h * 0.3, w * 0.5, -8)
    piece(s, "assets/collage/fern-sage.png", x + w * 0.7, y + h * 0.3, w * 0.45, 190)
    piece(s, CUT + "fungal-hyphae-1.png", x + w * 0.5, y + h * 0.68, w * 0.55, 90)


def farms():
    import scrapbook as sb
    d = Deck(name="02-10-2026-fri-ig-1000-farms-study")
    STK = " Stickers are drawn PowerPoint shapes (grouped, editable)."

    # 1 headline on a dark note, field collage above
    s = cream(d, "Template 2A, scrapbook. " + STUDY_SRC + " Field collage: stand-in from existing cutouts (ferns over torn-paper soil with the "
                 "fungal-hyphae cutout) until field-cross-section is generated [COLLAGE NEEDED]." + STK)
    tl = sb.note(s, 110, 70, 860, 560, sb.PAPER, 2)
    field(s, 150, 110, 780, 480)
    sb.seedling(s, 170, 640, 170, deg=-6)
    tl = sb.note(s, 60, 690, 960, 520, DEEP, -1.5)
    block(s, 110, 745, 860, "“A new study of working farms: the most regenerative farms stored 39% more carbon in their soil.”",
          52, CREAM, HEAD, True, 1.12)
    sb.beetle(s, 980, 700, 110, 25)
    source_line(s, STUDY_LINE)
    logo_br(s, False, 100)

    # 2 who measured
    s = cream(d, "Photo: assets/photo/erc-rancho-cacachilas-agro.jpg (Foundation library)." + STK)
    head_body(s, "The 1000 Farms Initiative, led by Ecdysis Foundation,",
              "measured farms across North America, from the most conventional to the most regenerative.", hpt=58, bpt=42)
    x, y, fw, fh = sb.polaroid(s, "assets/photo/erc-rancho-cacachilas-agro.jpg", 260, 600, 520, -3, "Foundation library")
    sb.tape(s, x + 60, y + 8, deg=-30); sb.tape(s, x + fw - 60, y + 8, deg=30)
    sb.magnifier(s, 890, 800, 90)
    sb.seedling(s, 150, 1170, 200, deg=-4)
    source_line(s, STUDY_LINE)

    # 3 and 4 big numbers
    def big(num, rest, note, stickers):
        s = cream(d, "Template 5 big number, scrapbook. " + note + STK)
        tl = sb.note(s, 110, 170, 860, 640, sb.PAPER, -2)
        text(s, 110, 230, 860, 300, num, 220, GREEN, HEAD, True, align="c", anchor="m", spacing=1.0, tl=tl)
        block(s, 170, 540, 740, rest, 50, DEEP, HEAD, True, 1.12, align="c")
        stickers(s)
        source_line(s, STUDY_LINE)
    big("39%", "more total soil carbon on the most regenerative farms.", "Photo: assets/photo/hand-soil-roots-fungi.jpg (Foundation library).",
        lambda s: (sb.polaroid(s, "assets/photo/hand-soil-roots-fungi.jpg", 90, 850, 330, -5, "Foundation library", 13),
                   sb.seedling(s, 620, 1180, 230), sb.mushroom(s, 860, 1180, 200, 6)))
    big("77%", "more total fungi in regenerative soils.", "Cutout: fungal-hyphae-1 (cut paper).",
        lambda s: (sb.cutout(s, CUT + "fungal-hyphae-1.png", 300, 1050, 420, 80), sb.mushroom(s, 700, 1190, 230, -4),
                   sb.mushroom(s, 900, 1180, 170, 8, cap="C9A227")))

    # 5 more life everywhere
    s = cream(d, "Stickers: bacteria, beetle, wildflowers and a bird, one for each group in the sentence." + STK)
    head_body(s, "More life everywhere: soil microbes, insects, plants and birds.", "The more biodiversity, the more carbon stored.",
              hpt=58, bpt=44)
    sb.bird(s, 780, 690, 260, -4)
    sb.flower(s, 560, 800, 60); sb.flower(s, 700, 900, 48, 10, petal="9DB8D9"); sb.flower(s, 880, 870, 55, -8)
    sb.beetle(s, 330, 1010, 150, -20)
    sb.bacterium(s, 170, 780, 170, -12); sb.bacterium(s, 260, 1210, 150, 18, "D6E3A8")
    sb.seedling(s, 620, 1230, 200); sb.mushroom(s, 850, 1230, 170, 6)
    source_line(s, STUDY_LINE)

    # 6 yields and profit
    s = cream(d, "Photo: assets/photo/erc-rancho-cacachilas-agro8.jpg (Foundation library)." + STK)
    head_body(s, "Regenerative yields matched national averages,", "and net profit per acre was similar.", hpt=60, bpt=46)
    x, y, fw, fh = sb.polaroid(s, "assets/photo/erc-rancho-cacachilas-agro8.jpg", 90, 560, 500, -4, "Foundation library")
    sb.tape(s, x + fw / 2, y + 6, 160, 44, 4)
    sb.coins(s, 820, 820, 190, -6); sb.coins(s, 900, 1030, 150, 8)
    sb.seedling(s, 740, 1230, 220)
    source_line(s, STUDY_LINE)

    # 7 the whole system
    s = cream(d, "One lonely seedling on the left, the whole system (fungi, bacteria, beetle, flowers) on the right." + STK)
    head_body(s, "One practice alone changed nothing.",
              "Farms using a single regenerative practice looked like conventional farms. The whole system matters.", hpt=60, bpt=44)
    tl = sb.note(s, 70, 700, 330, 460, sb.PAPER, -3)
    sb.seedling(s, 235, 1060, 220, sad=True)
    tl = sb.note(s, 450, 660, 560, 520, sb.PAPER, 2)
    sb.cutout(s, CUT + "fungal-hyphae-1.png", 730, 1030, 300, 90)
    sb.seedling(s, 640, 1010, 190); sb.flower(s, 860, 830, 50); sb.mushroom(s, 900, 1110, 150)
    sb.bacterium(s, 600, 780, 130, -10); sb.beetle(s, 760, 800, 90, 20)
    source_line(s, STUDY_LINE)

    # 8 read the paper
    s = cream(d, "Closing. Paper title: not reachable from this session (doi.org and the journal were blocked); "
                 "fill in the exact title from the DOI page [PAPER TITLE NEEDED]." + STK)
    tl = sb.note(s, 60, 120, 960, 780, DEEP, -1)
    y = block(s, 110, 180, 860, "Read the paper: Lundgren et al., Environmental Research: Food Systems, 2026.", 58, CREAM, HEAD, True, 1.12) + 50
    rect(s, 110, y, 140, 8, GLOW); y += 40
    text(s, 110, y, 860, 40, "[PAPER TITLE NEEDED]", 24, "E36B5E", BODY, True); y += 55
    text(s, 110, y, 860, 40, "Environmental Research: Food Systems", 24, GLOW, BODY); y += 50
    text(s, 110, y, 860, 40, "https://doi.org/10.1088/2976-601X/ae8f4e", 24, GLOW, BODY)
    sb.paper_doc(s, 300, 1090, 200, -8); sb.magnifier(s, 420, 1060, 80)
    sb.cutout(s, CUT + "dried-flowers-1-1.png", 700, 1120, 200, 18)
    source_line(s, STUDY_LINE)
    logo_br(s, False, 110)
    save(d, "02-10-2026-fri-ig-1000-farms-study.pptx")


def farms_li():
    import scrapbook as sb
    d = Deck(1200, 627, name="02-10-2026-fri-li-1000-farms-study")
    s = d.slide(DEEP, "LinkedIn image. " + STUDY_SRC + " Right half: the slide 1 field collage on a taped paper card (not tilted), "
                      "on Food Web Green.", counter=False)
    rect(s, 600, 0, 600, 627, GREEN)
    sb.note(s, 650, 60, 500, 430, sb.PAPER, 0)
    field(s, 670, 80, 460, 390)
    sb.seedling(s, 690, 570, 110, deg=-6)
    hh = 2 * lh(64, 1.05); sh = 2 * lh(24, 1.25); top = (627 - (hh + 24 + sh)) / 2
    text(s, 60, top, 520, hh + 10, "39% more\nsoil carbon", 64, CREAM, HEAD, True, spacing=1.05)
    text(s, 60, top + hh + 24, 520, sh + 10, "on the most regenerative farms,\n1000 Farms Initiative, 2026", 24, GLOW, BODY, spacing=1.25)
    logo_br(s, True, 90, 1200, 627, 30)
    save(d, "02-10-2026-fri-li-1000-farms-study.pptx")


# ================================================================ 5. World Teachers' Day trading cards (Mon 5 Oct)
MENTORS = [
    # name, role, advice, photo, (fx, fy), facts (based, since, background, known for)
    ("Tommy Tepper", "Director of Education & Mentor", None, None, None, ["[VERIFY]"] * 4),
    ("Loida Vasquez", "Advanced Programs Lead & Mentor", None, "assets/photo/loida-teaching-3.jpg", (0.3, 0.2), ["[VERIFY]"] * 4),
    ("Dr. Carla Portugal", "Science Lead & Mentor",
     "“Bare soil erodes. Living roots and cover hold the aggregates together.”", MENT + "carla-portugal.jpg", (0.5, 0.25),
     ["Brazil [VERIFY]", "2019", "PhD in Environmental Sciences, 20 years of environmental and farm consulting", "[VERIFY]"]),
    ("Wesley Sanders", "AP Mentor",
     "“Compost can look finished and still lack the biology you need. Check it under the microscope before you apply it.”",
     MENT + "wes-sander.jpg", (0.5, 0.35),
     ["Sierra Nevada foothills, California", "2020", "10 years as an agricultural journalist, then 10 years managing a farm",
      "Foothill Biological Soil Health Services"]),
    ("Casey Williams", "AP Mentor", None, None, None, ["[VERIFY]"] * 4),
    ("Brian Daubenspeck", "AP Mentor", None, None, None, ["[VERIFY]"] * 4),
    ("Isadora Schmidt", "AP Mentor", None, None, None, ["[VERIFY]"] * 4),
    ("Aysen Ustunay", "AP Mentor", None, None, None, ["[VERIFY]"] * 4),
    ("Dora Tkalec", "AP Mentor", None, None, None, ["[VERIFY]"] * 4),
    ("Gerald Ramirez", "AP Mentor",
     "“Extracts pull organisms off the compost into solution, so you can apply biology across a whole field.”",
     MENT + "gerald-ramirez.jpg", (0.5, 0.3),
     ["Costa Rica", "[VERIFY]", "Agronomist, University of Costa Rica", "[VERIFY]"]),
    ("Elena Kalli", "AP Admin", None, None, None, ["[VERIFY]"] * 4),
    ("Ib Borup Pederson", "AP Mentor", None, None, None, ["[VERIFY]"] * 4),
    ("Nick Padwick", "AP Mentor", "“I make 750 tons of compost a year. Biology works at any scale.”", MENT + "Nick.png", (0.5, 0.06),
     ["West Norfolk, England", "[VERIFY]", "Farmers Weekly Farm Manager and Farmer of the Year, 2009",
      "Managing Ken Hill Estate, home of Wild Ken Hill"]),
    ("Delvin Solkinson", "Permaculture Lead Teacher", None, None, None, ["[VERIFY]"] * 4),
]
LEAF_FRAME = CUT + "leaf-frame-1.png"   # not generated (Replicate throttled); leaf_frame() builds one from existing pieces


def leaf_frame(s, flip=False):
    """Leaf frame from cutouts already in the library: ferns, leaf sprigs and dried flowers around the edges."""
    P = [("assets/collage/fern-green.png", 300, 200, 520, -20), ("assets/collage/fern-sage.png", 800, 190, 460, 200),
         ("assets/collage/leaf-sprig-tan.png", 110, 640, 170, 4), ("assets/collage/leaf-sprig-blue.png", 975, 700, 160, -6),
         (CUT + "dried-flowers-1-1.png", 230, 1110, 330, -35), (CUT + "dried-flowers-2-1.png", 860, 1120, 280, 30),
         ("assets/collage/fern-green.png", 560, 1215, 420, 175)]
    for rel, cx, cy, w, deg in P:
        if flip: cy = 1350 - cy; deg = -deg
        piece(s, rel, cx, cy, w, deg)


class _G:
    """Lets lib drawing helpers draw into a group shape."""
    def __init__(self, grp, deck): self.shapes, self._deck = grp.shapes, deck


def card(d, name, role, advice, photo, focus, facts):
    note = (f"Trading card: {name}. Facts from the mentor bios collected on 30 Sep 2026 (soilfoodweb.com was blocked from this session); "
            "confirm every fact and the advice line with the mentor before posting.")
    if photo is None: note += f" [PHOTO NEEDED: {name}, headshot]"
    if photo and "carla" in photo: note += " Carla's headshot is only 499 px; ask for a larger file."
    if name == "Gerald Ramirez": note += " Bio also says 'Teaching the CLP and CTP courses' (left off: no acronyms in public text)."
    s = dark(d, note)
    grp = s.shapes.add_group_shape(); g = _G(grp, d)
    X, Y, W, H = 90, 55, 900, 1240            # card master: every card uses exactly these numbers
    rrect(g, X, Y, W, H, CREAM, radius=28, shadow=True)
    if photo:
        photo_box(g, photo, X + 36, Y + 36, W - 72, 520, *focus, radius=24)
    else:
        logo_box(g, X + 36, Y + 36, W - 72, 520, radius=24)
    npt = fit(name, W - 72, [56, 52, 48], 1, True)
    text(g, X + 36, Y + 570, W - 72, 80, name, npt, DEEP, HEAD, True, anchor="m")
    text(g, X + 36, Y + 652, W - 72, 40, role, 26, GREEN, BODY, True)
    adv = advice or "[ADVICE LINE NEEDED]"
    text(g, X + 36, Y + 705, W - 72, 260, adv, 34, DEEP if advice else RED, HEAD, True, spacing=1.12)
    labels = ("BASED IN", "TEACHING SINCE", "BACKGROUND", "KNOWN FOR")
    for i, (lab, val) in enumerate(zip(labels, facts)):
        cx = X + 36 + (i % 2) * 380; cy = Y + 985 + (i // 2) * 118
        text(g, cx, cy, 350, 26, lab, 15, GREEN, HEAD, True, track=2)
        text(g, cx, cy + 26, 350, 90, val, 19, INK, BODY, spacing=1.05)
    logo(g, X + W - 36 - 90, Y + H - 36 - 81, 90, white=False)
    grp.rotation = 0   # cards are never tilted
    return s


def paper(s, name, cx, cy, w, deg=0, fallback=None):
    """Cut-paper piece from assets/collage/cutouts/<name>-1.png; the drawn sticker if it was never generated."""
    import scrapbook as sb
    rel = CUT + name + "-1.png"
    if os.path.exists(os.path.join(ROOT, rel)):
        return sb.cutout(s, rel, cx, cy, w, deg)
    MISSING.add(rel)
    if fallback: return fallback()


CARD_STICKERS = {  # one cut-paper piece per card, loosely tied to what the mentor teaches
    "Dr. Carla Portugal": "happy-seedling", "Wesley Sanders": "microscope-paper", "Gerald Ramirez": "cute-bacterium",
    "Nick Padwick": "mushroom-paper", "Loida Vasquez": "earthworm-paper", "Delvin Solkinson": "leaf-sprig-paper",
}
ROTATION = ["cute-bacterium", "leaf-sprig-paper", "earthworm-paper", "mushroom-paper", "happy-seedling"]


def card_extras(s, name, i, has_photo):
    """Washi tape on the card corners, one cut-paper piece on the card edge, a leaf sprig on empty photo boxes."""
    import scrapbook as sb
    sb.tape(s, 150, 80, 170, 50, -32); sb.tape(s, 930, 80, 170, 50, 32)
    kind = CARD_STICKERS.get(name, ROTATION[i % len(ROTATION)])
    paper(s, kind, 985, 945, 170, 8, lambda: sb.seedling(s, 985, 1000, 150, deg=8))
    if not has_photo:
        paper(s, "leaf-sprig-paper", 860, 560, 170, 28, lambda: sb.cutout(s, CUT + "dried-flowers-2-1.png", 880, 560, 150, 24))


def teachers():
    import scrapbook as sb
    d = Deck(name="05-10-2026-mon-ig-teachers-day")
    src = "Source: mentor roster from Linnea, September 2026, and mentor bios on https://soilfoodweb.com. "
    frame = os.path.exists(os.path.join(ROOT, CUT + "leaf-frame-1.png"))
    s = dark(d, "Opening slide. " + src + "Cut-paper pieces from tools/collage_generate.py (flux-2-pro).")
    if frame: sb.cutout(s, CUT + "leaf-frame-1.png", 540, 690, 1060)
    else: leaf_frame(s)
    tl = sb.note(s, 170, 340, 740, 640, sb.PAPER, 0)
    block(s, 220, 410, 620, "Happy World Teachers' Day. Our mentors' best advice, in one line each.", 50, DEEP, HEAD, True, 1.12, align="c")
    paper(s, "microscope-paper", 830, 1080, 220, 6, lambda: sb.microscope(s, 800, 1060, 170, 6))
    paper(s, "earthworm-paper", 260, 1080, 190, -8, lambda: sb.seedling(s, 290, 1040, 150, deg=-6))
    logo_br(s, True)
    for i, m in enumerate(MENTORS):
        s = card(d, *m)
        card_extras(s, m[0], i, m[3] is not None)
    s = dark(d, "Closing slide. Cut-paper pieces from tools/collage_generate.py (flux-2-pro).")
    if frame: sb.cutout(s, CUT + "leaf-frame-1.png", 540, 690, 1060, 180)
    else: leaf_frame(s, flip=True)
    tl = sb.note(s, 190, 400, 700, 540, sb.PAPER, 0)
    block(s, 230, 480, 620, "Thank you to every mentor who teaches our students to see soil.", 50, DEEP, HEAD, True, 1.12, align="c")
    paper(s, "happy-seedling", 300, 1010, 190, -6, lambda: sb.flower(s, 300, 1000, 45, -8))
    paper(s, "cute-bacterium", 790, 1010, 210, -10, lambda: sb.bacterium(s, 790, 1020, 150, -12))
    logo_br(s, True)
    save(d, "05-10-2026-mon-ig-teachers-day.pptx")


# ================================================================ 6. World Habitat Day story (Mon 5 Oct), template 7
def habitat():
    import scrapbook as sb
    d = Deck(1080, 1920, name="05-10-2026-mon-ig-story-habitat-day")
    s = cream(d, "Story, scrapbook. Source: Anthony, Bender and van der Heijden, PNAS, 2023, https://doi.org/10.1073/pnas.2304663120. "
                 "Microscopy print: assets/microscopy/testate-amoeba-encysting-40x-joy-kaluf.jpg (Joy Kaluf, per the file name). "
                 "Leave y 1100 to 1450 empty for the poll sticker. Stickers are drawn shapes plus the nematode and fungal-hyphae cutouts.",
              1080, 1920)
    x, y, fw, fh = sb.polaroid(s, "assets/microscopy/testate-amoeba-encysting-40x-joy-kaluf.jpg", 250, 150, 560, -3, "Photo: Joy Kaluf", 15)
    sb.tape(s, x + fw / 2, y + 6, 170, 46, 4)
    sb.cutout(s, CUT + "nematode-1.png", 130, 330, 200, -30)
    sb.bacterium(s, 950, 300, 170, -15); sb.beetle(s, 960, 560, 110, 20)
    y = block(s, 100, 780, 880, "The biggest habitat on Earth is under your feet.", 70, DEEP, HEAD, True, 1.08, align="c")
    block(s, 100, y + 30, 880, "More than half of all species live in soil.", 44, INK, BODY, False, 1.2, align="c")
    sb.cutout(s, CUT + "fungal-hyphae-1.png", 180, 1720, 300, 100)
    sb.seedling(s, 850, 1780, 200); sb.mushroom(s, 560, 1790, 170, -4); sb.flower(s, 330, 1650, 45, 8)
    logo(s, (1080 - 180) / 2, 1500, 180, white=False)
    save(d, "05-10-2026-mon-ig-story-habitat-day.pptx")


# ================================================================ 7 to 10. Meet our mentors (single image, reused layout)
def mentor(d, photo, name, role, focus=(0.5, 0.3), note="", corner="br", sticker="seedling"):
    """Scrapbook: the photo as a big taped print on the top 65%, a Deep Green panel below with the name."""
    import scrapbook as sb
    s = cream(d, note + " Stickers are drawn shapes.")
    tl = Tilt(540, 450, -2)
    rect(s, 50, 40, 980, 820, sb.PAPER, tl=tl, shadow=True)
    if photo: pp = crop(photo, 940, 780, *focus); sh = s.shapes.add_picture(pp, 0, 0, Emu(1), Emu(1)); _place(sh, 70, 60, 940, 780, tl)
    else:
        rect(s, 70, 60, 940, 780, CREAM, tl=tl, pattern=STRIPE)
        logo(s, 540 - 210, 450 - 190, 420, white=False)
    sb.tape(s, 120, 60, 180, 52, -30); sb.tape(s, 960, 60, 180, 52, 30)
    rect(s, 0, 878, 1080, 472, DEEP)
    sprig = CUT + "dried-flowers-2-1.png"            # the same sprig used in the post 5 leaf frame
    if corner == "br": piece(s, sprig, 985, 810, 190, 28)
    else: piece(s, sprig, 95, 130, 190, -28)
    {"seedling": lambda: sb.seedling(s, 930, 1160, 190), "microscope": lambda: sb.microscope(s, 930, 1175, 200, 4),
     "cup": lambda: sb.cup(s, 900, 1080, 170, 6), "mushroom": lambda: sb.mushroom(s, 930, 1160, 190, 6),
     "flower": lambda: sb.flower(s, 920, 1000, 55, 8)}[sticker]()
    rrect(s, 80, 920, 330, 60, GLOW, radius=30)
    text(s, 80, 920, 330, 60, "Meet our mentors", 24, DEEP, HEAD, True, align="c", anchor="m")
    npt = fit(name, 760, [60, 54, 50], 1, True)
    text(s, 80, 1010, 800, 90, name, npt, CREAM, HEAD, True, anchor="m")
    text(s, 80, 1105, 760, 50, role, 30, GLOW, BODY)
    logo_br(s, True)
    return s


def carla():
    d = Deck(name="06-10-2026-tue-ig-mentor-carla-portugal")
    mentor(d, MENT + "carla-portugal.jpg", "Dr. Carla Portugal", "Science Lead & Mentor", (0.5, 0.2),
           "First 'Meet our mentors' post; posts 8 to 10 reuse this layout. No photo of Carla at a microscope in the library; "
           "this is her headshot (499 px, upscaled: ask for a larger file or a microscope photo) [PHOTO NEEDED: higher resolution].",
           sticker="seedling")
    save(d, "06-10-2026-tue-ig-mentor-carla-portugal.pptx")


def nick():
    d = Deck(name="07-10-2026-wed-ig-mentor-nick-padwick")
    mentor(d, MENT + "Nick.png", "Nick Padwick", "AP Mentor", (0.5, 0.1),
           "Photo: assets/mentors-teachers-day/Nick.png (Nick in a field with a spade). No photo with his windrows or compost turner "
           "in the library; the Drive photo from Wild Soils, Nov 2024, is the same small group shot as assets/photo/carla-nicks-son-nick-eri-wild-soils-event-11-2024.jpg.",
           corner="tl", sticker="mushroom")
    save(d, "07-10-2026-wed-ig-mentor-nick-padwick.pptx")


def wes():
    d = Deck(name="08-10-2026-thu-ig-mentor-wesley-sanders")
    mentor(d, MENT + "wes-sander.jpg", "Wesley Sanders", "AP Mentor", (0.5, 0.38),
           "Photo: assets/mentors-teachers-day/wes-sander.jpg (headshot outdoors). No photo of Wes at his microscope and no soil comparison "
           "video of his in assets/video/, so no still was used [PHOTO NEEDED: Wes at his microscope]. "
           "Name as given in the brief; the website and earlier posts spell it 'Wes Sander' [VERIFY].", sticker="microscope")
    save(d, "08-10-2026-thu-ig-mentor-wesley-sanders.pptx")


def gerald_caterina():
    d = Deck(name="09-10-2026-fri-ig-mentors-gerald-caterina")
    mentor(d, MENT + "gerald-teaching.jpg", "Gerald Ramirez", "AP Mentor", (0.22, 0.4),
           "Photo: assets/mentors-teachers-day/gerald-teaching.jpg (Gerald teaching a workshop). Confirm it is the Costa Rica workshop, "
           "March 2025, demonstrating liquid amendments [VERIFY]. Earlier posts spell his name Ramírez [VERIFY].", sticker="cup")
    mentor(d, None, "Dr. Caterina Capri", "Advanced Programs Instructor", note="[PHOTO NEEDED: Caterina headshot from Allison]",
           corner="tl", sticker="flower")
    save(d, "09-10-2026-fri-ig-mentors-gerald-caterina.pptx")


# ================================================================ 11. Soil Regenerators in the wild: Lisa Price (Sat 10 Oct), template 1
def lisa():
    import scrapbook as sb
    d = Deck(name="10-10-2026-sat-ig-soil-regenerators-lisa-price")
    s = cream(d, "Template 1A, scrapbook. [PHOTO NEEDED: Lisa with her seedlings] Photo box holds the colour logo until then. "
                 "Seedling tray (vegetable, herb and flower seedlings in little pots) is a drawn sticker.")
    tl = Tilt(540, 560, 2)
    rect(s, 90, 150, 900, 820, sb.PAPER, tl=tl, shadow=True)
    rect(s, 115, 175, 850, 700, CREAM, tl=tl, pattern=STRIPE)
    logo(s, 540 - 200, 525 - 180, 400, white=False)
    sb.tape(s, 150, 160, 180, 52, -30); sb.tape(s, 930, 160, 180, 52, 30)
    rrect(s, 80, 60, 620, 70, DEEP, radius=35)
    text(s, 80, 60, 620, 70, "Soil Regenerators in the wild", 26, GLOW, HEAD, True, align="c", anchor="m")
    sb.seedling_tray(s, 800, 1010, 380, -4)
    tl = sb.note(s, 60, 1080, 960, 200, DEEP, -1)
    npt = max(p for p in (64, 60, 56, 52) if text_width("Lisa Price, Australia", p) <= 780 or p == 52)
    text(s, 110, 1130, 800, 100, "Lisa Price, Australia", npt, CREAM, HEAD, True, anchor="m", tl=tl)
    logo(s, 900, 1150, 80, white=True)
    sb.cutout(s, CUT + "dried-flowers-1-1.png", 120, 1040, 190, -25)
    save(d, "10-10-2026-sat-ig-soil-regenerators-lisa-price.pptx")


POSTS = {1: coffee, 2: farms, 3: farms_li, 5: teachers, 6: habitat, 7: carla, 8: nick, 9: wes, 10: gerald_caterina, 11: lisa}

if __name__ == "__main__":
    for n in [int(a) for a in sys.argv[1:]] or POSTS:
        POSTS[n]()
    if MISSING: print("MISSING cutouts:", ", ".join(sorted(MISSING)))
