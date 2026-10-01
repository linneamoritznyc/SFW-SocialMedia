"""Boubacar carousel, root-network style, revised after Stephanie's feedback (30 Sep 2026):
- One fixed grid: the photo area and the text card sit in exactly the same place on every slide
  (main photo left, a narrow right column, a text card of fixed size and fixed distance from the bottom).
- A bolder, busier hyphae background that is one continuous panorama: the strands cross every seam and the colour
  blends gold, green, brown from slide to slide, so the 11 slides side by side make one design.
Brand fonts: Montserrat headings, Source Sans 3 body, EB Garamond italic for Boubacar's own words.

python3 tools/make_hyphae_panorama.py assets/collage/boubacar-panorama gold green brown gold green brown gold green brown gold green
python3 build_boubacar_panorama.py  ->  01-10-2026-thu-ig-soil-regenerators-boubacar-panorama.pptx
"""
import os
from PIL import ImageFont
from pptx.util import Emu
from lib import Deck, rect, text, crop, ROOT, PX, GOLD
from build_oct_01_10 import logo, save
from build_boubacar import B, SOURCES

W, H = 1080, 1350
HEADF, BODYF, VOICE = "Montserrat", "Source Sans 3", "EB Garamond"
PANO = "assets/collage/boubacar-panorama/slide-{:02d}.jpg"
ORDER = ["gold", "green", "brown", "gold", "green", "brown", "gold", "green", "brown", "gold", "green"]
CARD = {"gold": ("F6E3C2", "1E1412"), "green": ("E4E6CF", "22371F"), "brown": ("F1E4D6", "3A2524")}
FRAME = {"gold": "F6E3C2", "green": "E4E6CF", "brown": "F1E4D6"}
FOOT = {"gold": "1E1412", "green": "F4F1EA", "brown": "F4F1EA"}

# ---- the grid (identical on every slide) ----
PHOTO_Y, PHOTO_H = 80, 720                   # photo area: y 80 to 800
MAIN_X, MAIN_W = 80, 580                     # main photo, left
COL_X, COL_W = 720, 280                      # right column
CARD_X, CARD_W, CARD_Y, CARD_H = 60, 960, 860, 380   # text card: y 860 to 1240, always
PAD = 50
FONTFILE = {HEADF: "~/.fonts/Montserrat-Bold.ttf", BODYF: "~/.fonts/SourceSans3-Regular.ttf", VOICE: "~/.fonts/EBGaramond-Italic.ttf"}
_F = {}


def tw(t, pt, font):
    k = (font, pt)
    if k not in _F: _F[k] = ImageFont.truetype(os.path.expanduser(FONTFILE[font]), int(pt * 1.3333))
    return _F[k].getlength(t)


def nlines(t, pt, w, font):
    n = 0
    for para in t.split("\n"):
        n += 1; cur = ""
        for word in para.split(" "):
            trial = (cur + " " + word).strip()
            if cur and tw(trial, pt, font) > w: n += 1; cur = word
            else: cur = trial
    return n


def lh(pt, sp): return pt * 1.3333 * 1.17 * sp


def slide(d, i, note):
    s = d.slide(None, note, counter=False)
    s.shapes.add_picture(os.path.join(ROOT, PANO.format(i)), 0, 0, Emu(W * PX), Emu(H * PX))
    return s


def pic(s, theme, rel, x, y, w, h, fx=0.5, fy=0.5):
    rect(s, x - 10, y - 10, w + 20, h + 20, FRAME[theme], shadow=True)
    s.shapes.add_picture(crop(rel, int(w), int(h), fx, fy), Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))


def main(s, theme, rel, fx=0.5, fy=0.5):
    pic(s, theme, rel, MAIN_X, PHOTO_Y, MAIN_W, PHOTO_H, fx, fy)


def col1(s, theme, rel, fx=0.5, fy=0.5):
    pic(s, theme, rel, COL_X, PHOTO_Y, COL_W, PHOTO_H, fx, fy)


def col2(s, theme, a, b, fa=(0.5, 0.5), fb=(0.5, 0.5)):
    h = (PHOTO_H - 40) / 2
    pic(s, theme, a, COL_X, PHOTO_Y, COL_W, h, *fa)
    pic(s, theme, b, COL_X, PHOTO_Y + h + 40, COL_W, h, *fb)


def card(s, theme, body, head=None, voice=False):
    """Fixed-size card. Sizes step down until the text fits; the text is centred vertically in the card."""
    bg, fg = CARD[theme]
    rect(s, CARD_X, CARD_Y, CARD_W, CARD_H, bg, alpha=94, shadow=True)
    inner_w, inner_h = CARD_W - 2 * PAD, CARD_H - 2 * PAD + 10
    bfont = VOICE if voice else BODYF
    for hpt, bpt in ((52, 54), (50, 50), (48, 46), (46, 42), (44, 40), (42, 38), (40, 36), (38, 34), (36, 32), (34, 30)):
        hh = nlines(head, hpt, inner_w * 0.96, HEADF) * lh(hpt, 1.05) + 16 if head else 0
        bh = nlines(body, bpt, inner_w * 0.96, bfont) * lh(bpt, 1.15)
        if hh + bh <= inner_h: break
    y = CARD_Y + (CARD_H - hh - bh) / 2
    if head:
        text(s, CARD_X + PAD, y, inner_w, hh, head, hpt, fg, HEADF, True, spacing=1.05); y += hh
    text(s, CARD_X + PAD, y, inner_w, bh + 10, body, bpt, fg, bfont, False, italic=voice, spacing=1.15)


def quote_mark(s, theme, x=COL_X + 10, y=PHOTO_Y - 30, pt=260):
    text(s, x, y, COL_W, 300, "“", pt, "F6E3C2" if theme != "gold" else "1E1412", VOICE, False, spacing=0.8)


def footer(s, theme, t="soilfoodweb.com"):
    text(s, 0, H - 92, W, 50, t, 24, FOOT[theme], HEADF, True, align="c", anchor="m")


def build():
    d = Deck(name="01-10-2026-thu-ig-soil-regenerators-boubacar-panorama")
    N = (" Grid and continuous hyphae panorama after Stephanie's feedback. Background: assets/collage/boubacar-panorama/ "
         "(tools/make_hyphae_panorama.py, drawn by code). His quotes from his email to Allison, 30 Sep 2026.")

    s = slide(d, 1, "Cover." + N + " " + SOURCES + " Photo: boubacar-portrait-bananas-closer.jpg.")
    main(s, "gold", B + "boubacar-portrait-bananas-closer.jpg", 0.5, 0.35)
    logo(s, COL_X + 20, PHOTO_Y + 10, 240, white=False)
    text(s, COL_X, PHOTO_Y + 300, COL_W, 300, "Happy\nInternational\nCoffee Day", 28, "1E1412", HEADF, True, spacing=1.15)
    card(s, "gold", "Guinea-Conakry", "Boubacar Tidiane Diallo", voice=True)

    s = slide(d, 2, "Gnaly." + N + " Photos: farm-overview-with-tanks.jpeg, young-tree-with-pineapples.jpg, man-by-water-tanks.jpeg.")
    main(s, "green", B + "farm-overview-with-tanks.jpeg", 0.5, 0.5)
    col2(s, "green", B + "young-tree-with-pineapples.jpg", B + "man-by-water-tanks.jpeg", (0.5, 0.55), (0.35, 0.5))
    card(s, "green", "Fouta Djallon, Guinea-Conakry", "Gnaly Coffee & AgroÉcole Bio")
    footer(s, "green")

    s = slide(d, 3, "The question." + N + " Photo: vegetable-beds-by-building.jpeg.")
    main(s, "brown", B + "vegetable-beds-by-building.jpeg", 0.6, 0.25)
    quote_mark(s, "brown")
    card(s, "brown", "“Today, I ask a different question: What does the soil food web need to thrive?”", voice=True)
    footer(s, "brown")

    s = slide(d, 4, "Feeding the soil." + N + " Photos: boubacar-compost-pile.jpeg; still-biochar-kiln-smoke-73s.png "
                    "(making-biochar-in-barrel-kiln.mp4 at 1:13).")
    main(s, "gold", B + "boubacar-compost-pile.jpeg", 0.5, 0.45)
    col1(s, "gold", B + "still-biochar-kiln-smoke-73s.png", 0.5, 0.4)
    card(s, "gold", "living groundcover, mulch from his farm and the forest, and compost with biochar activated with fish "
                    "hydrolysate and compost tea.", "How he feeds his soil:")
    footer(s, "gold")

    s = slide(d, 5, "In practice." + N + " Photos: carrying-mulch.jpg, still-biochar-charcoal-16s.png, barrel-brew-under-shelter.jpeg.")
    main(s, "green", B + "carrying-mulch.jpg", 0.5, 0.5)
    col2(s, "green", B + "still-biochar-charcoal-16s.png", B + "barrel-brew-under-shelter.jpeg", (0.5, 0.5), (0.5, 0.35))
    card(s, "green", "“I apply compost regularly around young trees and whenever I add new mulch.”", voice=True)
    footer(s, "green")

    s = slide(d, 6, "What grows." + N + " Photos: planting-seedling-in-agroforest.jpeg, banana-bunch.jpg.")
    main(s, "brown", B + "planting-seedling-in-agroforest.jpeg", 0.45, 0.45)
    col1(s, "brown", B + "banana-bunch.jpg", 0.5, 0.4)
    card(s, "brown", "coffee, cacao, bananas, citrus, and pineapple.", "What he grows together:")
    footer(s, "brown")

    s = slide(d, 7, "His compost." + N + " Photos: buckets-of-tubers.jpg, young-tree-yellow-new-leaves.jpg, man-with-harvested-tubers.jpeg [VERIFY who].")
    main(s, "gold", B + "buckets-of-tubers.jpg", 0.5, 0.5)
    col2(s, "gold", B + "young-tree-yellow-new-leaves.jpg", B + "man-with-harvested-tubers.jpeg", (0.5, 0.45), (0.5, 0.5))
    card(s, "gold", "“My compost includes plant residues, animal manure when available, biochar activated with fish hydrolysate "
                    "and compost tea, and local organic materials.”", voice=True)
    footer(s, "gold")

    s = slide(d, 8, "What's changing." + N + " Photos: millipede-leaf-litter.jpg, boubacar-hand-compost-worm.jpg.")
    main(s, "green", B + "millipede-leaf-litter.jpg", 0.5, 0.35)
    col1(s, "green", B + "boubacar-hand-compost-worm.jpg", 0.4, 0.5)
    card(s, "green", "more earthworms, more mushrooms, active decomposition, and better soil structure.", "What he's seeing:")
    footer(s, "green")

    s = slide(d, 9, "Life returns." + N + " Photos: yellow-caterpillar-on-stem.jpg, caterpillar-on-leaf.jpg, still-young-plants-buckets-29s.png.")
    main(s, "brown", B + "yellow-caterpillar-on-stem.jpg", 0.35, 0.4)
    col2(s, "brown", B + "caterpillar-on-leaf.jpg", B + "still-young-plants-buckets-29s.png", (0.5, 0.35), (0.5, 0.55))
    card(s, "brown", "“I can already see signs of healthier soil in the field, including more earthworms, mushrooms, and active "
                     "decomposition.”", voice=True)
    footer(s, "brown")

    s = slide(d, 10, "Passing it on." + N + " Photos: group-of-five-farmers.jpeg, two-men-at-shade-house.jpeg [VERIFY who, where, when].")
    main(s, "gold", B + "group-of-five-farmers.jpeg", 0.42, 0.5)
    col1(s, "gold", B + "two-men-at-shade-house.jpeg", 0.5, 0.35)
    card(s, "gold", "“Your training changed the way I see agriculture. I now believe that agriculture is biology, not chemical "
                    "fertilizer.”", voice=True)
    footer(s, "gold")

    s = slide(d, 11, "The dream and the call to action (YouTube)." + N + " Photo: young-coffee-plant.jpg.")
    main(s, "green", B + "young-coffee-plant.jpg", 0.5, 0.5)
    quote_mark(s, "green")
    logo(s, COL_X + 20, PHOTO_Y + PHOTO_H - 216, 240, white=True)
    card(s, "green", "“My dream is to restore degraded land in the Fouta Djallon and help my community learn living-soil practices.”",
         voice=True)
    footer(s, "green", "Subscribe on YouTube: @boubacartidianediallo")

    save(d, "01-10-2026-thu-ig-soil-regenerators-boubacar-panorama.pptx")


if __name__ == "__main__":
    build()
