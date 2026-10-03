"""World Animal Day carousel, version 2 (Sunday 4 October 2026, Instagram, 10 slides, 1080 x 1350).
Bold marketing style (like the 1000 Farms carousel): full-bleed photo on top, brand green panel below with cream type,
category labels with line icons, one source line per slide. The brief's copy is used word for word, split over 10 slides.
Photos: only the first version's (assets/photo/animals/). Icons: assets/icons/ (tools/icon_generate.py, flux-2-pro).

python3 build_world_animal_day_v2.py  ->  04-10-2026-sun-ig-world-animal-day-v2.pptx
Every element is a separate editable object.
"""
import os
from PIL import Image
from pptx.util import Emu, Pt
from lib import Deck, rect, text, crop, est_lines, ROOT, PX, HEAD, BODY, rgb

# Brand palette (docs/brand-colors.md, the Canva "2026 SFW Foundation Brand Main colors"); red only for the STAYS OUT label
BG, CREAM, BROWN, TAN, SAGE, GOLD = "31662F", "F3F1EA", "4C3634", "C09D7F", "B1BCB1", "D39C48"
from build_oct_01_10 import logo, save

W, H, M = 1080, 1350, 80
CW = W - 2 * M
RED = "B23A2E"
AN = "assets/photo/animals/"
ICON = "assets/icons/icon-{}-{}.png"
PH = 700                              # photo height on photo slides; Deep Green panel below

SOURCES = ("Sources: Soil Food Web School, Compost Manual: [LINK FROM CARLA]. "
           "Soil Food Web School, BioComplete Compost, Lecture 4: Thermophilic Composting (Part 2): https://pt.soilfoodweb.com/?p=2067. "
           "Miller et al. (2015). Ascariasis in humans and pigs on small-scale farms, Maine, USA, 2010 to 2013. "
           "Emerging Infectious Diseases, 21(2): https://doi.org/10.3201/eid2102.140048. "
           "USDA NRCS and Fairbanks Soil and Water Conservation District (2005). Composting Dog Waste: "
           "https://www.epa.gov/system/files/documents/2022-11/Composting-Dog-Waste-Booklet-Alaska.pdf. "
           "CDC (2025). About Toxoplasmosis: https://www.cdc.gov/toxoplasmosis/about/index.html. "
           "Course: https://school.soilfoodweb.com/bundles/advanced-biocomplete-compost-production")


def credits():
    out = {}
    for line in open(os.path.join(ROOT, AN, "credits.txt")):
        f = line.rstrip("\n").split("\t")
        out[f[0].rsplit(".", 1)[0]] = dict(name=f[1].replace("Photo by ", "").replace(" on Unsplash", ""),
                                            page=f[3].replace("Photo page: ", ""))
    return out
CRED = credits()


def cred_line(*keys):
    return "Photo: " + ", ".join(CRED[k]["name"] for k in keys) + " on Unsplash"


def cred_note(*keys):
    return " ".join(f"Photo {k}.jpg: {CRED[k]['name']} on Unsplash, {CRED[k]['page']}." for k in keys)


def pic(s, key, x, y, w, h, fx=0.5, fy=0.5):
    p = crop(AN + key + ".jpg", int(w), int(h), fx, fy)
    return s.shapes.add_picture(p, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))


def icon(s, name, colour, x, y, h):
    p = os.path.join(ROOT, ICON.format(name, colour))
    iw, ih = Image.open(p).size; w = h * iw / ih
    return s.shapes.add_picture(p, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX))), w


def label(s, y, t, fill, fg, ico):
    """Solid label: line icon, then the text in 18 pt Montserrat Bold caps, tracking 2."""
    tw = len(t) * 18 * 1.3333 * 0.74
    w = 18 + 26 + 12 + tw + 18
    rect(s, M, y, w, 48, fill)
    icon(s, ico, "brandcream" if fg == CREAM else "brown", M + 18, y + 11, 26)
    text(s, M + 56, y, tw + 18, 48, t, 18, fg, HEAD, True, anchor="m", track=2)


def head(s, y, t, pt=60, color=CREAM, w=CW):
    h = est_lines(t, pt, w * 0.97, True) * pt * 1.3333 * 1.08 + 8
    text(s, M, y, w, h, t, pt, color, HEAD, True, spacing=1.05)
    return y + h


def body(s, y, t, pt=30, color=CREAM, w=CW):
    h = est_lines(t, pt, w * 0.97) * pt * 1.3333 * 1.3 + 10
    text(s, M, y, w, h, t, pt, color, BODY, spacing=1.3)
    return y + h


def rich(s, y, runs, pt=30, color=CREAM, x=M, w=CW):
    """One paragraph, several runs: (text, bold). Bold runs are Montserrat Bold, the rest Source Sans 3."""
    h = est_lines("".join(t for t, _ in runs), pt, w * 0.95) * pt * 1.3333 * 1.3 + 10
    tb = s.shapes.add_textbox(0, 0, Emu(1), Emu(1)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.line_spacing = 1.3
    for t, b in runs:
        r = p.add_run(); r.text = t; f = r.font
        f.size = Pt(pt); f.name = HEAD if b else BODY; f.bold = b; f.color.rgb = rgb(color)
    tb.left, tb.top, tb.width, tb.height = Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX))
    return y + h


def foot(s, src):
    """Source line bottom left, white logo bottom right, on the Deep Green panel."""
    text(s, M, H - 70, CW - 140, 24, src, 14, CREAM, BODY, alpha=70)
    logo(s, W - M - 90, H - 60 - 90 * 668 / 743, 90, white=True)


def credit(s, t, y=PH + 12):
    text(s, M, y, CW, 24, t, 14, CREAM, BODY, alpha=70, align="r")


def photo_slide(d, note, key, fx=0.5, fy=0.5):
    s = d.slide(BG, note, counter=False)
    pic(s, key, 0, 0, W, PH, fx, fy)
    credit(s, cred_line(key))
    return s


def build():
    d = Deck(name="04-10-2026-sun-ig-world-animal-day-v2")

    # 1 cover
    s = d.slide(None, "Cover. " + SOURCES + " " + cred_note("cow"), counter=False)
    pic(s, "cow", 0, 0, W, H, 0.5, 0.3)
    rect(s, 0, H / 2, W, H / 2, BG, alpha=40)
    text(s, M, 735, CW, 350, "Whose manure belongs in your compost?", 80, CREAM, HEAD, True, spacing=1.02)
    text(s, M, 1140, CW, 50, "Happy World Animal Day, October 4.", 30, CREAM, BODY)
    text(s, M, H - 42, CW, 24, cred_line("cow"), 14, CREAM, BODY, alpha=75)
    logo(s, W - M - 90, H - M - 90 * 668 / 743, 90, white=True)

    # 2 manure: a high-nitrogen feedstock
    s = d.slide(BG, "Feedstocks, text.", counter=False)
    icon(s, "manure-fork", "brandcream", M, 150, 150)
    y = head(s, 360, "Manure: a high-nitrogen feedstock", 64) + 30
    body(s, y, "A Soil Food Web thermophilic compost recipe combines three feedstocks: high-nitrogen material such as manure "
               "(about 10% in summer), green plant material (about 30%) and woody material (about 60%). Bedding, such as straw "
               "or wood shavings, counts toward the woody share.", 32)
    foot(s, "Soil Food Web School, BioComplete Compost, Lecture 4")

    # 3 the recipe as one bar: 10 / 30 / 60
    s = d.slide(BG, "Feedstocks, the recipe as one bar (widths 10, 30, 60 percent). Icons: assets/icons/.", counter=False)
    parts = [("manure-fork", "High nitrogen, 10%", TAN, 0.10), ("leaf", "Green, 30%", SAGE, 0.30), ("wood-log", "Woody, 60%", CREAM, 0.60)]
    for i, (ico, cap, col, share) in enumerate(parts):
        y = 170 + i * 300
        icon(s, ico, "brandcream", M, y, 110)
        text(s, M + 200, y + 20, CW - 200, 70, cap, 40, CREAM, HEAD, True, anchor="m")
        rect(s, M, y + 150, CW, 34, CREAM, alpha=15)
        rect(s, M, y + 150, CW * share, 34, col)
    foot(s, "Soil Food Web School, BioComplete Compost, Lecture 4")

    # 4 goes in: 4 x 2 grid, full bleed
    grid = ["rabbit", "goat", "sheep", "alpaca", "guinea-pig", "horse", "cow", "chicken"]
    focus = {"rabbit": (0.5, 0.45), "goat": (0.55, 0.5), "sheep": (0.35, 0.5), "alpaca": (0.5, 0.3), "guinea-pig": (0.45, 0.5),
             "horse": (0.5, 0.55), "cow": (0.5, 0.4), "chicken": (0.5, 0.5)}
    s = d.slide(BG, "Plant-eaters and poultry. " + cred_note(*grid), counter=False)
    g = 8; tw = (W - 3 * g) / 4
    for i, key in enumerate(grid):
        pic(s, key, (i % 4) * (tw + g), (i // 4) * (tw + g), tw, tw, *focus[key])
    gy = 2 * tw + g
    text(s, M, gy + 12, CW, 48, cred_line(*grid), 14, CREAM, BODY, alpha=70)
    label(s, gy + 76, "GOES IN", SAGE, BROWN, "check-mark")
    y = head(s, gy + 142, "Plant-eaters and poultry") + 8
    body(s, y, "Rabbit, goat, sheep, alpaca, guinea pig, horse, cow and chicken. Their manure is a great high-nitrogen "
               "ingredient. Use it fresh, or dry it and store it until you build the pile.")
    foot(s, "Soil Food Web School, Compost Manual")

    # 5 pig
    s = photo_slide(d, "Pig. " + cred_note("pig"), "pig", 0.6, 0.45)
    label(s, 760, "HOT PILE ONLY", GOLD, BROWN, "flame")
    y = head(s, 826, "Pig") + 8
    body(s, y, "Pig roundworm can infect people. Pig manure goes into a pile that completes the thermophilic requirements, "
               "and raw pig manure stays away from produce.")
    foot(s, "Miller et al. 2015, Emerging Infectious Diseases")

    # 6 dog (first half of the dog and cat text)
    s = photo_slide(d, "Dog and cat, part 1. " + cred_note("dog"), "dog", 0.5, 0.4)
    label(s, 760, "STAYS OUT", RED, CREAM, "crossed-circle")
    y = head(s, 826, "Dog and cat") + 8
    body(s, y, "Dog feces can carry roundworm (Toxocara) eggs. If swallowed, the larvae can migrate to the liver, lungs and eyes.")
    foot(s, "USDA NRCS 2005; CDC 2025")

    # 7 cat (second half)
    s = photo_slide(d, "Dog and cat, part 2. " + cred_note("cat"), "cat", 0.62, 0.6)
    label(s, 760, "STAYS OUT", RED, CREAM, "crossed-circle")
    body(s, 836, "Cat feces can carry Toxoplasma gondii, which people can also pick up from contaminated soil. USDA advises "
                 "against using dog waste compost on food crops, even after hot composting, and against composting cat waste "
                 "or litter at all.")
    foot(s, "USDA NRCS 2005; CDC 2025")

    # 8 slow or hot: cold pile
    s = photo_slide(d, "Slow or hot, part 1. " + cred_note("compost-2"), "compost-2", 0.5, 0.75)
    y = head(s, 760, "Slow or hot, both work.", 60) + 24
    rect(s, M, y, 6, 110, SAGE)
    rich(s, y, [("Cold or static pile: ", True), ("fine for plant-eater and poultry manure, on a longer timeline.", False)],
         x=M + 30, w=CW - 30)
    foot(s, "Soil Food Web School, Compost Manual")

    # 9 slow or hot: thermophilic pile
    s = photo_slide(d, "Slow or hot, part 2. Thermometer icon: assets/icons/. " + cred_note("compost-4"), "compost-4", 0.5, 0.6)
    icon(s, "compost-thermometer", "brandcream", W - M - 110, 760, 120)
    y = head(s, 760, "Slow or hot, both work.", 60, w=CW - 140) + 24
    rect(s, M, y, 6, 160, SAGE)
    rich(s, y, [("Thermophilic pile: ", True), ("above 55°C (131°F) for 3 days at the center. Turn it so every part passes "
                                                "through the hot center.", False)], x=M + 30, w=CW - 30)
    foot(s, "Soil Food Web School, Compost Manual")

    # 10 closing
    s = d.slide(BG, "Closing. Course: https://school.soilfoodweb.com/bundles/advanced-biocomplete-compost-production", counter=False)
    head(s, 170, "Advanced BioComplete Compost Production", 64)      # renders as four lines, to about y 600
    y = 640
    rect(s, M, y, 120, 4, TAN)
    body(s, y + 40, "Building and assessing three compost piles that meet biological minimums. At the Soil Food Web "
                    "School. Link in bio.", 32)   # course line from docs/copy-deck-v2.md
    lw = 240
    logo(s, (W - lw) / 2, 900 + (450 - lw * 668 / 743) / 2, lw, white=True)
    save(d, "04-10-2026-sun-ig-world-animal-day-v2.pptx")


if __name__ == "__main__":
    build()
