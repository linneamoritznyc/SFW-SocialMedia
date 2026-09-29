"""Teachers Day instructor carousel, 1080 x 1350, all native objects. Run: python3 build_teachers_day.py
Cover, four instructor cards (identical), closing. Card content is laid out flat then rotated about the card centre."""
import math, os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_PATTERN
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

PX = 9525; W, H = 1080, 1350; TILT = 3
BROWN, GREEN, DEEP, GOLD, CREAM, LIGHT, INK, WHITE = "4F3433", "156826", "22371F", "C9A227", "F4F1EA", "59A66C", "333130", "FFFFFF"
STRIPE, GOLD_D, GREEN_D = "E6E0D0", "9A7A14", "0E4419"
HEAD, BODY, TAG = "Montserrat", "Source Sans 3", "[COPY: Allison]"
prs = Presentation(); prs.slide_width = Emu(W*PX); prs.slide_height = Emu(H*PX)
rgb = RGBColor.from_string

def blank(bg, notes=""):
    s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(bg)
    if notes: s.notes_slide.notes_text_frame.text = notes
    return s

class Card:
    """Places shapes flat, then tilts each about the card centre."""
    def __init__(s_, s, cx, cy, tilt): s_.s, s_.cx, s_.cy, s_.a = s, cx, cy, math.radians(tilt); s_.t = tilt
    def _place(c, sh, x, y, w, h):
        dx, dy = x+w/2-c.cx, y+h/2-c.cy
        nx = c.cx + dx*math.cos(c.a) - dy*math.sin(c.a); ny = c.cy + dx*math.sin(c.a) + dy*math.cos(c.a)
        sh.left, sh.top, sh.width, sh.height = Emu(int((nx-w/2)*PX)), Emu(int((ny-h/2)*PX)), Emu(int(w*PX)), Emu(int(h*PX))
        sh.rotation = c.t
    def shape(c, kind, x, y, w, h, fill, line=None, lw=0, pattern=False, flip=False):
        r = c.s.shapes.add_shape(kind, 0, 0, Emu(1), Emu(1))
        if pattern:
            r.fill.patterned(); r.fill.pattern = MSO_PATTERN.LIGHT_UPWARD_DIAGONAL
            r.fill.fore_color.rgb = rgb(STRIPE); r.fill.back_color.rgb = rgb(fill)
        else: r.fill.solid(); r.fill.fore_color.rgb = rgb(fill)
        if line: r.line.color.rgb = rgb(line); r.line.width = Emu(int(lw*PX))
        else: r.line.fill.background()
        r.shadow.inherit = False; c._place(r, x, y, w, h)
        if flip: r._element.spPr.xfrm.set("flipH", "1")
        return r
    def label(c, sh, t, size, color, font, bold=True):
        tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = t
        r.font.size = Pt(size*0.75); r.font.name = font; r.font.bold = bold; r.font.color.rgb = rgb(color)
    def text(c, x, y, w, h, t, size, color, font, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = c.s.shapes.add_textbox(0, 0, Emu(1), Emu(1)); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.alignment = align; r = p.add_run(); r.text = t
        r.font.size = Pt(size*0.75); r.font.name = font; r.font.bold = bold; r.font.color.rgb = rgb(color)
        c._place(tb, x, y, w, h); return tb

def flat_text(s, x, y, w, h, t, size, color, font, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    return Card(s, 0, 0, 0).text(x, y, w, h, t, size, color, font, bold, align, anchor)

def title_slide(bg, color, head, sub, handle=None, notes=""):
    s = blank(bg, notes)
    flat_text(s, 100, 330, 880, 420, f"{TAG} {head}", 96, color, HEAD, True, anchor=MSO_ANCHOR.BOTTOM)
    flat_text(s, 100, 800, 880, 200, f"{TAG} {sub}", 44, color, BODY)
    if handle: flat_text(s, 100, 1140, 880, 60, handle, 36, GREEN, HEAD, True)

def instructor(i):
    s = blank(DEEP, "Instructor card. Photo: head and shoulders, smiling, written consent. Bio: use only what the instructor approved. No invented years or facts.")
    c = Card(s, 540, 675, TILT); R = MSO_SHAPE
    c.shape(R.RECTANGLE, 110, 130, 880, 1130, "0F1C0E")                      # shadow
    c.shape(R.RECTANGLE, 100, 110, 880, 1130, CREAM, pattern=True)           # card
    ph = c.shape(R.RECTANGLE, 140, 150, 340, 340, "A7B097", CREAM, 8)        # photo, cream border
    c.label(ph, "PHOTO", 40, DEEP, HEAD)
    bx, by, br = 780, 320, 130                                                 # scalloped badge
    for k in range(18):
        a = 2*math.pi*k/18; c.shape(R.OVAL, bx+math.cos(a)*(br-14)-28, by+math.sin(a)*(br-14)-28, 56, 56, CREAM, "D9D2BF", 2)
    c.shape(R.OVAL, bx-br+14, by-br+14, 2*br-28, 2*br-28, CREAM)
    lg = c.shape(R.OVAL, bx-92, by-92, 184, 184, WHITE, GOLD, 3); c.label(lg, "LOGO", 30, DEEP, HEAD)
    c.text(140, 540, 800, 300, f"{TAG} Bio, three short lines.\nYears teaching, what they teach, one fun detail.", 52, INK, BODY)
    # ribbons: fold tails behind, then bands
    c.shape(R.RECTANGLE, 128, 920, 60, 130, GOLD_D); c.shape(R.RECTANGLE, 892, 920, 60, 130, GOLD_D)
    c.shape(R.RECTANGLE, 168, 900, 744, 130, GOLD)
    nm = c.text(168, 900, 744, 130, f"{TAG} NAME", 56, WHITE, HEAD, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    c.shape(R.RECTANGLE, 188, 1020, 50, 96, GREEN_D); c.shape(R.RECTANGLE, 842, 1020, 50, 96, GREEN_D)
    c.shape(R.RECTANGLE, 220, 1010, 640, 96, GREEN)
    c.text(220, 1010, 640, 96, f"{TAG} Soil Food Web Instructor", 32, WHITE, HEAD, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    c.shape(R.RECTANGLE, 260, 1122, 560, 44, BROWN)
    c.text(260, 1122, 560, 44, "soilfoodweb.com", 32, CREAM, BODY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

title_slide(GOLD, DEEP, "Happy Teachers Day", "One short line: who we are celebrating.", notes="Cover.")
for i in range(4): instructor(i)
title_slide(CREAM, DEEP, "Thank you to all our instructors", "@soilfoodwebschool", notes="Closing.")
prs.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), "teachers-day-instructors.pptx"))
