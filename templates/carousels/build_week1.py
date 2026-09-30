"""Week 1 (5 to 11 Oct 2026): Mon 5 Oct Teachers' Day five mentors (7 slides), Tue 6 Oct Carla Portugal (2 slides).
Run: python3 build_week1.py   Cream background, Food Web Green circle frames, Deep Green Source Sans 3 tips, Montserrat Bold names."""
import os
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE
from photos import crop, ROOT
from PIL import Image, ImageDraw, ImageOps

MD = os.path.join(ROOT, "assets", "mentors-teachers-day")
def find(slug):
    for ext in ("jpg", "jpeg", "png"):
        f = os.path.join(MD, f"{slug}.{ext}")
        if os.path.exists(f): return f

def circle_png(path, d):
    im = ImageOps.fit(ImageOps.exif_transpose(Image.open(path)).convert("RGB"), (d, d), Image.LANCZOS)
    m = Image.new("L", (d, d), 0); ImageDraw.Draw(m).ellipse((0, 0, d, d), fill=255)
    im.putalpha(m); out = os.path.join(MD, f".circle-{d}.png"); im.save(out); return out
import build as B
from build import new, slide, rect, text, copy, handle, dots, C, W, H, PX, DISPLAY, SANS, TAG

def frame_photo(s, cx, y, d, label, slug=None):
    rect(s, cx-d/2-16, y-16, d+32, d+32, fill="green", shape=MSO_SHAPE.OVAL)       # Food Web Green ring
    f = find(slug) if slug else None
    if f:
        s.shapes.add_picture(circle_png(f, d), Emu(int((cx-d/2)*PX)), Emu(int(y*PX)), Emu(d*PX), Emu(d*PX)); return
    rect(s, cx-d/2, y, d, d, fill="sage", shape=MSO_SHAPE.OVAL)
    text(s, cx-d/2+40, y, d-80, d, f"[PHOTO: {label}]", 34, color="moss", font=DISPLAY, bold=True, align="c", anchor="m")

MENTORS = [
    ("Dr. Carla Portugal", "Science Leader", "carla-portugal"),
    ("Nick Padwick", "Farmer, Norfolk", "nick-padwick"),
    ("Wes Sander", "Microscopy", "wes-sander"),
    ("Gerald Ramirez", "Compost extracts and teas", "gerald-ramirez"),
    ("Dr. Caterina Capri", "Advanced Programs", "caterina-capri"),
]

def five_mentors():
    p = new()
    s = slide(p, "cream", "Slide 1 of 7. Group photo of mentors at a workshop (Costa Rica or Wild Ken Hill). Placeholder photo: assets/photo/workshop-group-around-compost-pile.jpg. Swap for a real mentor group photo.")
    s.shapes.add_picture(crop(os.path.relpath(find("group-mentors-workshop") or os.path.join(ROOT, "assets/photo/workshop-group-around-compost-pile.jpg"), ROOT), W, 800, 0.5, 0.5), 0, 0, Emu(W*PX), Emu(800*PX))
    rect(s, 0, 800, W, 550, fill="moss")
    text(s, 80, 850, 920, 60, "WORLD TEACHERS' DAY", 40, color="glow", font=DISPLAY, bold=True)
    text(s, 80, 930, 920, 330, "Five mentors.\nFive tips for anyone starting out.", 76, color="white", font=DISPLAY, bold=True)
    handle(s, "sage"); dots(s, 7, 0, on="glow", off="scope")

    for i, (name, role, slug) in enumerate(MENTORS):
        s = slide(p, "cream", f"Slide {i+2} of 7. {name}, {role}. Portrait: round headshot from the website, with consent. Tip: from Allison's question 'What is the one thing you wish every new student knew?' One or two sentences, in the mentor's words.")
        frame_photo(s, W/2, 110, 440, f"{name}", slug)
        text(s, 80, 620, 920, 80, name, 64, color="moss", font=DISPLAY, bold=True, align="c")
        text(s, 80, 705, 920, 60, role, 44, color="green", font=SANS, bold=False, align="c")
        copy(s, 80, 800, 920, 380, "“[PLACEHOLDER tip]”", 64, color="moss", font=SANS, align="c")
        handle(s); dots(s, 7, i+1)

    s = slide(p, "moss", "Slide 7 of 7. Photo of a mentor and student at a microscope. Real photo, with consent.")
    f = find("mentor-student-microscope")
    if f: s.shapes.add_picture(crop(os.path.relpath(f, ROOT), W, 760), 0, 0, Emu(W*PX), Emu(760*PX))
    else:
        rect(s, 0, 0, W, 760, fill="sage")
        text(s, 100, 0, 880, 760, "[PHOTO: a mentor and student at a microscope]", 38, color="moss", font=DISPLAY, bold=True, align="c", anchor="m")
    copy(s, 80, 830, 920, 300, "Thank you to every mentor. \U0001F49A", 76, color="white", font=DISPLAY, bold=True)
    handle(s, "sage"); dots(s, 7, 6, on="glow", off="scope")
    return p

def rounded_photo(s, x, y, w, h, slug, label):
    f = find(slug)
    if f:
        pic = s.shapes.add_picture(crop(os.path.relpath(f, ROOT), w, h), Emu(x*PX), Emu(y*PX), Emu(w*PX), Emu(h*PX))
        pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE; return
    rect(s, x, y, w, h, fill="sage", shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.08
    text(s, x+60, y, w-120, h, f"[PHOTO: {label}]", 38, color="moss", font=DISPLAY, bold=True, align="c", anchor="m")

def carla():
    p = new()
    s = slide(p, "cream", "Slide 1 of 2. Photo: Carla at a microscope or teaching at a workshop. Large rounded frame. Consent needed.")
    rounded_photo(s, 100, 100, 880, 700, "carla-teaching", "Carla at a microscope or teaching")
    text(s, 100, 860, 880, 60, "MEET OUR MENTORS", 44, color="gold", font=DISPLAY, bold=True)
    text(s, 100, 935, 880, 190, "Dr. Carla Portugal", 84, color="green", font=DISPLAY, bold=True)
    text(s, 100, 1120, 880, 60, "Science Leader", 48, color="moss", font=SANS)
    handle(s); dots(s, 2, 0)

    s = slide(p, "cream", "Slide 2 of 2. Tip comes from Monday's question. Image: a microscope image from Carla if she has a favorite.")
    rounded_photo(s, 100, 100, 880, 520, "carla-favorite-microscope", "Carla's favorite microscope image")
    text(s, 100, 680, 880, 60, "CARLA'S TIP", 44, color="green", font=DISPLAY, bold=True)
    copy(s, 100, 760, 880, 420, "“[PLACEHOLDER, from Monday's question]”", 64, color="moss", font=SANS)
    handle(s); dots(s, 2, 1)
    return p

if __name__ == "__main__":
    os.chdir(os.path.dirname(os.path.abspath(__file__)))
    os.makedirs("posts", exist_ok=True)
    for f in ("posts/2026-10-05-mon-did-you-know.pptx",):
        if os.path.exists(f): os.remove(f)
    a = five_mentors(); B.date_deck(a, __import__("datetime").date(2026,10,5), "World Teachers' Day: five mentors, five tips"); a.save("posts/2026-10-05-mon-teachers-day-five-mentors.pptx")
    c = carla(); B.date_deck(c, __import__("datetime").date(2026,10,6), "Mentor week: Dr. Carla Portugal"); c.save("posts/2026-10-06-tue-carla-portugal.pptx")
