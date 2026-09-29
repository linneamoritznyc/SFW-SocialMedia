"""Photo-led carousel, 1080 x 1350, all native objects. Run: python3 build_photo_led.py
Full-bleed PHOTO placeholder, rounded text panels, deep green footer bar with LOGO placeholder."""
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

PX = 9525; W, H = 1080, 1350
DEEP, GREEN, CREAM, GLOW, SAGE, INK = "22371F", "156826", "F4F1EA", "DBE6A7", "A7B097", "333130"
HEAD, BODY, TAG = "Montserrat", "Source Sans 3", "[COPY: Allison]"
prs = Presentation(); prs.slide_width = Emu(W*PX); prs.slide_height = Emu(H*PX)
rgb = RGBColor.from_string

def box(s, kind, x, y, w, h, fill, alpha=None, adj=None):
    r = s.shapes.add_shape(kind, Emu(x*PX), Emu(y*PX), Emu(w*PX), Emu(h*PX))
    r.fill.solid(); r.fill.fore_color.rgb = rgb(fill); r.line.fill.background(); r.shadow.inherit = False
    if alpha is not None:
        clr = r.fill._xPr.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
        a = etree.SubElement(clr, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha"); a.set("val", str(alpha*1000))
    if adj is not None: r.adjustments[0] = adj
    return r

def words(r, t, size, color, font, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, pad=30):
    tf = r.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = pad; tf.margin_top = tf.margin_bottom = 14
    p = tf.paragraphs[0]; p.alignment = align; p.line_spacing = 1.05
    run = p.add_run(); run.text = t; f = run.font; f.size = Pt(size*0.75); f.name = font; f.bold = bold; f.color.rgb = rgb(color)

def panel(s, y, h, t, style="deep", size=54):
    fill, alpha, col = {"deep": (DEEP, 88, CREAM), "cream": (CREAM, 92, DEEP), "green": (GREEN, 92, WHITE)}[style]
    r = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, 60, y, 960, h, fill, alpha, 0.12)
    words(r, f"{TAG} {t}", size, col, BODY)

WHITE = "FFFFFF"
def slide(panels, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ph = box(s, MSO_SHAPE.RECTANGLE, 0, 0, W, 1230, SAGE)
    words(ph, "PHOTO", 44, DEEP, HEAD)
    for p in panels: panel(s, *p)
    box(s, MSO_SHAPE.RECTANGLE, 0, 1230, W, 120, DEEP)
    lg = box(s, MSO_SHAPE.OVAL, 80, 1250, 80, 80, CREAM); words(lg, "LOGO", 18, DEEP, HEAD, pad=0)
    ft = s.shapes.add_textbox(Emu(180*PX), Emu(1250*PX), Emu(420*PX), Emu(80*PX)); tf = ft.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; r = tf.paragraphs[0].add_run(); r.text = "Soil Food Web School"
    r.font.size = Pt(34*0.75); r.font.name = HEAD; r.font.bold = True; r.font.color.rgb = rgb(GLOW)
    hd = s.shapes.add_textbox(Emu(580*PX), Emu(1250*PX), Emu(420*PX), Emu(80*PX)); tf = hd.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = "@soilfoodwebschool"; r.font.size = Pt(30*0.75); r.font.name = BODY; r.font.color.rgb = rgb(CREAM)
    s.notes_slide.notes_text_frame.text = notes
    return s

slide([(280, 110, "Happy [Occasion]!", "deep", 60), (470, 200, "One short line that celebrates it.", "green", 50)],
      "Cover. Photo: real SFW photo, people looking away or at the land works best. Keep both panels off faces.")
slide([(230, 250, "Fact one, two or three lines, plain words.", "deep", 52), (540, 250, "What follows from it, one to three lines.", "green", 52)],
      "Any fact needs a named source before posting. Do not paste unsourced science claims.")
slide([(200, 130, "Fact, one line.", "deep", 52), (350, 240, "What it means, two lines.", "green", 52)],
      "Keep panels above the crowd. Move a panel if it covers faces.")
slide([(320, 210, "Fact, two lines.", "deep", 52), (560, 130, "What it means, one line.", "green", 52)],
      "Same layout, panels lower.")
slide([(170, 230, "One closing call to action, two lines at most.", "deep", 56)],
      "Closing. One call to action only.")
prs.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), "photo-led-carousel.pptx"))
